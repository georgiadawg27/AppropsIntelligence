"""
Load-test the Labor-HHS CSVs the way they reach the workbook: a copy of the
reference workbook with each lhhs_*.csv merged into its tab -- a row whose ID
is already there (v30 on carries them) kept as the workbook has it (the
owner's workbook is the record: v32 updated SRC-CRPT-119HRPT271's
also_covers), any other appended; the
Component tab and the headline_observation_id column added if missing --
loaded by approps_store.load into a scratch database. Prints the load report,
the store's own warnings and the in-store Title II reconciliation.

    python reference/review/lhhs/workbook/merge_check.py [workbook]

Nothing is written to the reference workbook.
"""

import csv
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
sys.path.insert(0, str(ROOT))

import approps_store as S  # noqa: E402

TABS = {"Account": "account", "Historical Name": "historical_name", "Source Document": "source_document",
        "Bill Report Reference": "bill_report_reference", "Appropriations Observation": "observation",
        "Confirmed Absence": "confirmed_absence", "Account Relationship": "account_relationship",
        "Validation Record": "validation_record", "Component": "component"}


def merged(workbook, out):
    import openpyxl
    # values, not formulas: v28's bill_id / URL columns are lookups whose cached
    # values a re-saved copy wouldn't keep
    wb = openpyxl.load_workbook(workbook, data_only=True)
    for tab, name in TABS.items():
        rows = list(csv.DictReader(open(HERE / f"lhhs_{name}.csv", newline="")))
        if tab not in wb.sheetnames:
            ws = wb.create_sheet(tab)
            ws.append(list(rows[0]) if rows else [])
        ws = wb[tab]
        head = [c.value for c in ws[1]]
        if rows:
            for col in rows[0]:
                if col not in head:                      # headline_observation_id: the proposed column
                    ws.cell(row=1, column=len(head) + 1, value=col)
                    head.append(col)
        # an empty template row (v28's Account Relationship tab has one) would be read as data
        if ws.max_row == 2 and all(c.value is None for c in ws[2]):
            ws.delete_rows(2)
        # a row already in the workbook under the same ID (column A) is the workbook's,
        # never duplicated or overwritten: v30 on already carries these rows, v28 didn't
        have = {str(ws.cell(row=i, column=1).value) for i in range(2, ws.max_row + 1)}
        for r in rows:
            if r[head[0]] not in have:
                ws.append([convert(r.get(h)) for h in head])
    wb.save(out)


def convert(v):
    if v is None or v == "":
        return None
    for cast in (int, float):
        try:
            return cast(v)
        except ValueError:
            pass
    return v


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    workbook = Path(argv[0]) if argv else S.reference_workbook()
    with tempfile.TemporaryDirectory() as tmp:
        wb = Path(tmp) / "v29_candidate.xlsx"
        merged(workbook, wb)
        report = S.load(wb, Path(tmp) / "approps.db")
        conn = S.connect(Path(tmp) / "approps.db", readonly=True)
        report["title_ii"] = title_ii_in_store(conn)
        conn.close()
        print("rows:", report["rows"])
        print("tabs not in workbook:", report["tabs_not_in_workbook"], "| value map:", report["value_map"])
        print(f"{len(report['warnings'])} warnings")
        for w in report["warnings"]:
            print("  ", w)
        ok = sum(r["reconciles_through_rollups"] == "yes" for r in report["title_ii"])
        print(f"Title II total reconciles through the stored rollups in {ok} of {len(report['title_ii'])} cells")
        for r in report["title_ii"]:
            print(f"   FY{r['fiscal_year']} {r['stage']}: {r['reconciles_through_rollups']}"
                  + (f" (differs by {r['through_rollups_differs_by_thousands']})" if r["reconciles_through_rollups"] != "yes" else ""))
        return report


def title_ii_in_store(conn):
    """Each recorded Title II total, reconciled from the store alone
    (title_totals.reconcile, through the rollups: the agency totals, AHA in the
    requests, the General Provisions accounts; CURES and the rescissions as
    signed candidates)."""
    sys.path.insert(0, str(ROOT / "reference" / "review"))
    import title_totals as T
    printed = [{"fiscal_year": r["fiscal_year"], "stage": r["stage"],
                "printed_total_title_iii_thousands": r["amount"] // 1000}
               for r in conn.execute("SELECT fiscal_year, stage, amount FROM appropriations_observation WHERE "
                                     "canonical_account_id = 'ACC-HHS-TITLE-II-TOTAL' AND component IS NULL "
                                     "ORDER BY fiscal_year, stage")]
    return T.reconcile(conn, printed, title="Title II", subcommittee="LHHS")


if __name__ == "__main__":
    main()
