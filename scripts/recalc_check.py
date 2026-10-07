"""
Recalculate a built workbook in LibreOffice (headless) and fail on any formula error.

    python3 scripts/recalc_check.py build/Approps_Pilot_Schema_Loaded.xlsx [--expect-formulas N]

Opens the workbook through LibreOffice's UNO bridge, recalculates every formula
(calculateAll), then walks every used cell of every sheet:

  - a formula cell whose result is an error (#N/A, #REF!, #VALUE!, Err:5xx, ...) fails the check;
  - a text cell whose text starts with "=" is counted as text: data that begins with "="
    is never a formula, so it must not appear among the formula cells.

With --compare-cached, every formula's result must also equal the result stored in the file
(a workbook whose results were filled in by approps_store.calculate_lookups, or saved by a
spreadsheet that recalculated): LibreOffice's calculation checks the stored one.

Prints one summary line (formulas, errors, text cells starting with "=") and exits 1 on
any error, a stored result that differs, or when --expect-formulas is given and the
formula count differs. Needs
LibreOffice and its Python bridge (Debian/Ubuntu: libreoffice-calc python3-uno, plus
python3-openpyxl for --compare-cached); run it with the system python3, which is the one
that can import uno.
"""

import argparse
import os
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import uno  # noqa: E402  (LibreOffice's Python bridge)
from com.sun.star.beans import PropertyValue  # noqa: E402
from com.sun.star.connection import NoConnectException  # noqa: E402

STRING_CELLS, FORMULA_CELLS = 4, 16   # com.sun.star.sheet.CellFlags.STRING / FORMULA
ERROR_RESULT = 4                      # com.sun.star.sheet.FormulaResult.ERROR


def prop(name, value):
    p = PropertyValue()
    p.Name, p.Value = name, value
    return p


def connect(port, profile):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise SystemExit("LibreOffice (soffice) not found")
    proc = subprocess.Popen([soffice, "--headless", "--invisible", "--nologo", "--norestore", "--nodefault",
                             f"-env:UserInstallation={Path(profile).as_uri()}",
                             f"--accept=socket,host=127.0.0.1,port={port};urp;"],
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    resolver = uno.getComponentContext().ServiceManager.createInstanceWithContext(
        "com.sun.star.bridge.UnoUrlResolver", uno.getComponentContext())
    for _ in range(120):
        try:
            ctx = resolver.resolve(f"uno:socket,host=127.0.0.1,port={port};urp;StarOffice.ComponentContext")
            return proc, ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        except NoConnectException:
            time.sleep(0.5)
    proc.kill()
    raise SystemExit("could not connect to LibreOffice")


def cached_results(path):
    """{(sheet, row, col) 0-based: stored result} for every formula cell, as the file holds them."""
    import openpyxl                                  # read with the same python only when asked
    wf = openpyxl.load_workbook(path, read_only=True)
    wv = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wf.worksheets:
        for r, (frow, vrow) in enumerate(zip(ws.iter_rows(), wv[ws.title].iter_rows(values_only=True))):
            for c, cell in enumerate(frow):
                if getattr(cell, "data_type", None) == "f":         # a formula; text that starts with "=" is not
                    out[(ws.title, r, c)] = vrow[c] if c < len(vrow) else None
    return out


def check(path, port=2002, cached=None):
    with tempfile.TemporaryDirectory() as profile:
        proc, desktop = connect(port, profile)
        try:
            doc = None
            for _ in range(20):             # the first load can come back empty while LibreOffice finishes starting
                doc = desktop.loadComponentFromURL(Path(path).resolve().as_uri(), "_blank", 0, (prop("Hidden", True),))
                if doc is not None:
                    break
                time.sleep(1)
            if doc is None:
                raise SystemExit(f"LibreOffice could not open {path}")
            doc.calculateAll()
            formulas, errors, eq_text, differ = 0, [], 0, []
            sheets = doc.getSheets()
            for i in range(sheets.getCount()):
                sh = sheets.getByIndex(i)
                for rng in sh.queryContentCells(FORMULA_CELLS).getRangeAddresses():
                    formulas += (rng.EndRow - rng.StartRow + 1) * (rng.EndColumn - rng.StartColumn + 1)
                    if cached is not None:
                        block = sh.getCellRangeByPosition(rng.StartColumn, rng.StartRow, rng.EndColumn, rng.EndRow)
                        for dr, row in enumerate(block.getDataArray()):
                            for dc, got in enumerate(row):
                                at = (sh.getName(), rng.StartRow + dr, rng.StartColumn + dc)
                                want = cached.get(at)
                                want = "" if want is None else str(want)
                                got = got if isinstance(got, str) else ("" if got is None else str(got))
                                if got != want:
                                    differ.append(f"{at[0]} row {at[1] + 1} col {at[2] + 1}: LibreOffice {got!r}, stored {want!r}")
                for rng in sh.queryFormulaCells(ERROR_RESULT).getRangeAddresses():
                    for r in range(rng.StartRow, rng.EndRow + 1):
                        for c in range(rng.StartColumn, rng.EndColumn + 1):
                            cell = sh.getCellByPosition(c, r)
                            errors.append(f"{sh.getName()}!{cell.AbsoluteName.split('.')[-1]}: error {cell.getError()}")
                for rng in sh.queryContentCells(STRING_CELLS).getRangeAddresses():
                    block = sh.getCellRangeByPosition(rng.StartColumn, rng.StartRow, rng.EndColumn, rng.EndRow)
                    eq_text += sum(1 for row in block.getDataArray() for v in row if isinstance(v, str) and v.startswith("="))
            doc.close(True)
        finally:
            try:
                desktop.terminate()
            except Exception:  # noqa: BLE001 -- LibreOffice exits as it answers
                pass
            proc.wait(timeout=60)
    return formulas, errors, eq_text, differ


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--expect-formulas", type=int)
    ap.add_argument("--compare-cached", action="store_true",
                    help="also require every formula's stored result to equal LibreOffice's")
    ap.add_argument("--port", type=int, default=int(os.environ.get("RECALC_PORT", "2002")))
    a = ap.parse_args(argv)
    cached = cached_results(a.workbook) if a.compare_cached else None
    formulas, errors, eq_text, differ = check(a.workbook, a.port, cached)
    print(f"LibreOffice recalculation: {formulas:,} formulas, {len(errors)} errors, "
          f"{eq_text} text cells starting with '=' kept as text ({formulas + eq_text:,} cells start with '=')")
    for e in errors[:50]:
        print("  " + e)
    if cached is not None:
        print(f"stored results: {len(cached) - len(differ):,} of {len(cached):,} equal LibreOffice's")
        for d in differ[:50]:
            print("  " + d)
    bad = bool(errors) or bool(differ)
    if a.expect_formulas is not None and formulas != a.expect_formulas:
        print(f"expected {a.expect_formulas:,} formulas, found {formulas:,}")
        bad = True
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
