"""
The relational store and its query (approps_store.py, store_schema.sql),
loaded from the committed pilot workbook (reference/..._v10.xlsx).

Run:  python -m unittest tests.test_store -v

Needs openpyxl (loading only).
"""

import contextlib
import io
import shutil
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import approps_store as S  # noqa: E402

try:
    import openpyxl
except ImportError:                                  # pragma: no cover
    openpyxl = None

WORKBOOK = ROOT / "reference" / "CJS_Title_III_Science_Pilot_Schema_Loaded_v10.xlsx"
FOUR = ["President's Budget", "House Reported", "Senate Reported", "Enacted"]


def workbook_rows(tab):
    """A tab read straight from the workbook -- independent of the loader."""
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    rows = list(wb[tab].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:] if r[0] is not None]


@unittest.skipUnless(openpyxl, "openpyxl not installed")
class StoreTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.db = Path(cls.tmp.name) / "approps.db"
        cls.report = S.load(WORKBOOK, cls.db)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def setUp(self):
        self.conn = S.connect(self.db)

    def tearDown(self):
        self.conn.close()

    def query(self, text, agency=None):
        res = S.resolve(self.conn, text, agency)
        return res, (S.history(self.conn, res["account"]["canonical_account_id"]) if res["account"] else None)


class Load(StoreTest):
    def test_all_seven_tabs_load(self):
        self.assertEqual(self.report["rows"], {
            "account": 30, "historical_name": 2, "source_document": 23, "bill_report_reference": 41,
            "appropriations_observation": 983, "account_relationship": 2, "validation_record": 89})

    def test_foreign_keys_are_enforced_not_just_declared(self):
        # SQLite ignores REFERENCES unless the connection turns them on
        with self.assertRaises(sqlite3.IntegrityError):
            self.conn.execute("INSERT INTO historical_name VALUES ('HN-X', 'ACC-NOPE', 'n', 'e', NULL, 1, 1)")
        self.assertEqual(self.conn.execute("PRAGMA foreign_key_check").fetchall(), [])

    def test_workbook_booleans_are_booleans(self):
        # 'TRUE' as text would evaluate false in SQLite
        want = sum(1 for o in workbook_rows("Appropriations Observation") if o["offsetting_collections"] == "TRUE")
        self.assertGreater(want, 0)
        got = self.conn.execute("SELECT count(*) FROM appropriations_observation WHERE offsetting_collections").fetchone()[0]
        self.assertEqual(got, want)
        with self.assertRaises(sqlite3.IntegrityError):              # STRICT: no text in an INTEGER column
            self.conn.execute("UPDATE appropriations_observation SET offsetting_collections = 'TRUE'")

    def test_spelling_variant_is_mapped_and_counted(self):
        self.assertEqual(self.report["value_map"],
                         {"Appropriations Observation.extraction_method: 'human_entered' -> 'human-entered'": 983})

    def test_dates_are_iso_dates(self):
        self.assertEqual(self.conn.execute("SELECT publication_date FROM source_document "
                                           "WHERE document_id = 'SRC-CRPT-119HRPT652'").fetchone()[0], "2026-05-15")

    def test_bill_report_reference_is_a_real_key(self):
        r = self.conn.execute("SELECT o.bill_id, o.report_id, b.bill_id, b.report_id FROM appropriations_observation o "
                              "JOIN bill_report_reference b ON b.reference_id = o.bill_report_reference_id "
                              "WHERE o.observation_id = 'OBS-0082'").fetchone()
        self.assertEqual(tuple(r), ("S.2354", "S.Rept.119-44", "S.2354", "S.Rept.119-44"))
        self.assertEqual(self.conn.execute("SELECT count(*) FROM appropriations_observation "
                                           "WHERE bill_report_reference_id IS NULL").fetchone()[0], 0)

    def test_document_identity_questions_are_warned(self):
        text = "\n".join(self.report["warnings"])
        self.assertIn("SRC-CRPT-119HRPT652: a committee_report recorded at stage 'Enacted'", text)
        self.assertIn("SRC-EXPL-FY2026-PB", text)


class NasaScienceAcceptance(StoreTest):
    def test_resolves_through_the_extraction_matching_rule(self):
        res, _ = self.query("NASA Science")
        self.assertEqual((res["account"]["canonical_account_id"], res["match"], res["via"], res["agency"]),
                         ("ACC-NASA-SCIENCE", "exact", "canonical", "National Aeronautics and Space Administration"))
        # the same letters-only edit distance and threshold as extraction
        self.assertEqual(self.query("NASA Scince")[0]["match"], "ocr_corrected")

    def test_history_equals_the_workbook_row_for_row(self):
        _, h = self.query("NASA Science")
        got = sorted((o["observation_id"], o["fiscal_year"], o["stage"], o["amount_type"], o["amount"],
                      o["source_document_id"], o["source_page"]) for o in h["observations"])
        want = sorted((o["observation_id"], o["fiscal_year"], o["stage"], o["amount_type"], o["amount"],
                       o["source_document_id"], o["source_page"])
                      for o in workbook_rows("Appropriations Observation")
                      if o["canonical_account_id"] == "ACC-NASA-SCIENCE")
        self.assertEqual(got, want)
        self.assertEqual(len(got), 40)
        self.assertEqual({(o["fiscal_year"], o["stage"]) for o in h["observations"]},
                         {(y, s) for y in range(2017, 2027) for s in FOUR})

    def test_fy2027_is_reported_missing_not_zero(self):
        _, h = self.query("NASA Science")
        self.assertEqual(h["missing_cells"], [(2027, s) for s in FOUR])

    def test_values_match_the_printed_documents(self):
        # read off the documents themselves, not the workbook: text layer
        # (Senate), page image by eye (House 115-231 p.120, 116-101 p.148),
        # recorded vision transcription (House 119-652 p.169)
        _, h = self.query("NASA Science")
        cell = {(o["fiscal_year"], o["stage"]): o for o in h["observations"]}
        for (fy, stage), amount, doc, page in (
                ((2024, "Senate Reported"), 7_340_920_000, "SRC-CRPT-118SRPT62", 219),
                ((2024, "President's Budget"), 8_260_800_000, "SRC-CRPT-118SRPT62", 219),
                ((2023, "Enacted"), 7_795_000_000, "SRC-CRPT-118SRPT62", 219),
                ((2026, "Senate Reported"), 7_300_000_000, "SRC-CRPT-119SRPT44", 217),
                ((2026, "Enacted"), 7_250_000_000, "SRC-CRPT-119HRPT652", 169),
                ((2020, "House Reported"), 7_161_300_000, "SRC-CRPT-116HRPT101", 148),
                ((2017, "Enacted"), 5_764_900_000, "SRC-CRPT-115HRPT231", 120),
                ((2018, "President's Budget"), 5_711_800_000, "SRC-CRPT-115HRPT231", 120)):
            o = cell[(fy, stage)]
            lo, hi = map(int, o["source_page"].split("-"))
            self.assertEqual((o["amount"], o["source_document_id"]), (amount, doc), (fy, stage))
            self.assertTrue(lo <= page <= hi, (fy, stage, o["source_page"]))

    def test_every_value_cites_a_source_document_that_exists(self):
        _, h = self.query("NASA Science")
        docs = {r[0] for r in self.conn.execute("SELECT document_id FROM source_document")}
        self.assertTrue(all(o["source_document_id"] in docs and o["url_or_identifier"] for o in h["observations"]))
        # and that document is, or also covers, the observation's cell
        for o in h["observations"]:
            d = dict(self.conn.execute("SELECT * FROM source_document WHERE document_id = ?",
                                       (o["source_document_id"],)).fetchone())
            self.assertIn((o["fiscal_year"], o["stage"]), S.covered_cells(d)[0], o["observation_id"])


class ExplorationAcceptance(StoreTest):
    def test_base_and_supplemental_split(self):
        _, h = self.query("NASA Exploration")
        obs = h["observations"]
        self.assertEqual(len(obs), 80)
        by = {(o["fiscal_year"], o["stage"], o["amount_type"]): o["amount"] for o in obs}
        self.assertEqual({t for (_, _, t) in by}, {"budget authority", "supplemental"})
        self.assertEqual(by[(2024, "Enacted", "budget authority")], 7_216_200_000)
        self.assertEqual(by[(2024, "Enacted", "supplemental")], 450_000_000)
        self.assertEqual(by[(2025, "Senate Reported", "supplemental")], 1_212_000_000)

    def test_former_name_reaches_the_same_account(self):
        for text, kind in (("Deep Space Exploration Systems", "exact"), ("Deep Spaee Exploratlon Systems", "ocr_corrected")):
            res, h = self.query(text)
            self.assertEqual((res["account"]["canonical_account_id"], res["match"], res["via"]),
                             ("ACC-NASA-EXPLORATION", kind, "historical_name"))
            self.assertEqual(len(h["observations"]), 80)

    def test_relationship_is_shown_not_merged(self):
        _, h = self.query("Exploration")
        self.assertEqual([(r["relationship_id"], r["other_account_id"]) for r in h["relationships"]],
                         [("REL-0001", "ACC-NASA-EXPLTECH")])
        self.assertNotIn("ACC-NASA-EXPLTECH", {o["canonical_account_id"] for o in h["observations"]})

    def test_matching_pool_reads_review_state_from_the_store(self):
        self.conn.execute("UPDATE historical_name SET human_reviewed = 0 WHERE historical_name_id = 'HN-0001'")
        try:
            self.assertIsNone(self.query("Deep Space Exploration Systems")[0]["account"])
        finally:
            self.conn.rollback()


class Resolve(StoreTest):
    def test_same_name_in_two_agencies_is_not_guessed(self):
        res, _ = self.query("Office of Inspector General")
        self.assertEqual((res["match"], res["account"]), ("ambiguous", None))
        self.assertEqual([c["canonical_account_id"] for c in res["candidates"]][:2], ["ACC-NASA-OIG", "ACC-NSF-OIG"])
        self.assertEqual(self.query("NSF Office of Inspector General")[0]["account"]["canonical_account_id"], "ACC-NSF-OIG")
        self.assertEqual(self.query("Office of Inspector General", agency="NASA")[0]["account"]["canonical_account_id"],
                         "ACC-NASA-OIG")

    def test_agency_alone_is_its_agency_total(self):
        self.assertEqual(self.query("NASA")[0]["account"]["canonical_account_id"], "ACC-NASA-TOTAL")

    def test_cli_exit_codes(self):
        for args, code in ((["NASA Science", "--json"], 0), (["Office of Inspector General"], 2),
                           (["Nothing Like Any Account"], 1)):
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                self.assertEqual(S.main(["history", *args, "--db", str(self.db)]), code, args)


class StoreCopyTest(StoreTest):
    """Tests that write: each gets its own copy of the loaded database."""

    def setUp(self):
        self.path = Path(self.tmp.name) / f"{self.id().rsplit('.', 1)[-1]}.db"
        shutil.copy(self.db, self.path)
        self.conn = S.connect(self.path)

    def flagged(self, account):
        return [r[0] for r in self.conn.execute(
            "SELECT observation_id FROM appropriations_observation WHERE canonical_account_id = ? "
            "AND verification_status = 'flagged' ORDER BY 1", (account,))]


class ResolutionCascade(StoreCopyTest):
    def test_v10_resolution_that_reached_one_observation_is_reported(self):
        text = "\n".join(self.report["warnings"])
        self.assertIn("validation_record VAL-0081 resolves ACC-NASA-EXPLTECH's identity, but its own observation is "
                      "still flagged (OBS-0522); 2 other observation(s) of the account are still flagged with no "
                      "resolved record: OBS-0523, OBS-0524", text)
        self.assertIn("validation_record VAL-0082 resolves ACC-NASA-LEO's identity", text)

    def test_resolution_reaches_every_observation_it_decides(self):
        s = S.resolve_relationship(self.conn, "REL-0001", "Reviewer", "withdrawn proposal, never adopted")
        self.assertEqual(s["observations"], ["OBS-0522", "OBS-0523", "OBS-0524"])
        self.assertEqual(s["records_updated"], ["VAL-0081"])
        self.assertEqual(s["records_created"], ["VAL-REL-0001-OBS-0523", "VAL-REL-0001-OBS-0524"])
        self.assertEqual(self.flagged("ACC-NASA-EXPLTECH"), [])
        for oid in s["observations"]:
            r = self.conn.execute("SELECT human_review_status, reviewer, resolution FROM validation_record "
                                  "WHERE observation_id = ? AND rule_applied = 'account_identity'", (oid,)).fetchall()
            self.assertEqual([tuple(x) for x in r], [("resolved", "Reviewer", "withdrawn proposal, never adopted")], oid)
        self.assertEqual(self.conn.execute("SELECT human_reviewed FROM account_relationship "
                                           "WHERE relationship_id = 'REL-0001'").fetchone()[0], 1)
        # the other relationship's observations are not this resolution's
        self.assertEqual(self.flagged("ACC-NASA-LEO"), ["OBS-0525", "OBS-0526", "OBS-0527"])
        self.assertFalse(any("ACC-NASA-EXPLTECH" in w for w in S.stale_resolution_warnings(self.conn)))

    def test_another_open_finding_keeps_its_observation_flagged(self):
        self.conn.execute("INSERT INTO validation_record (validation_id, observation_id, rule_applied, result) "
                          "VALUES ('VAL-X', 'OBS-0523', 'table_total', 'fail')")
        self.conn.commit()
        s = S.resolve_relationship(self.conn, "REL-0001", "Reviewer", "withdrawn proposal")
        self.assertEqual(s["still_flagged"], ["OBS-0523"])
        self.assertEqual(self.flagged("ACC-NASA-EXPLTECH"), ["OBS-0523"])

    def test_refused_resolution_changes_nothing(self):
        with self.assertRaises(LookupError):
            S.resolve_relationship(self.conn, "REL-9999", "Reviewer", "x")
        with self.assertRaises(ValueError):
            S.resolve_relationship(self.conn, "REL-0001", "", "x")
        self.assertEqual(self.flagged("ACC-NASA-EXPLTECH"), ["OBS-0522", "OBS-0523", "OBS-0524"])


HOUSE_PDF = ROOT / "document_store" / "CRPT-119hrpt652.pdf"
VISION_FIXTURES = ROOT / "tests" / "fixtures" / "vision_cache"


@unittest.skipUnless(HOUSE_PDF.exists(), "CRPT-119hrpt652.pdf not present")
class PipelineToStore(StoreCopyTest):
    """The vision path's Title III output goes into the store and answers
    the same query -- the extraction half and the store half joined."""

    def test_extraction_output_loads_and_is_queryable(self):
        import extract_approps as ex
        with tempfile.TemporaryDirectory() as out:
            r = ex.run(HOUSE_PDF, title="TITLE III", cache_dir=VISION_FIXTURES, offline=True, out_dir=out, verbose=False)
        # only this document's facts, so nothing is counted twice
        self.conn.execute("DELETE FROM validation_record")
        self.conn.execute("DELETE FROM appropriations_observation")
        rows = [row for row in (S.observation_row(o, "SRC-CRPT-119HRPT652") for o in r["observations"]) if row]
        with self.conn:
            for row in rows:
                self.conn.execute(f"INSERT INTO appropriations_observation ({', '.join(row)}) "
                                  f"VALUES ({', '.join('?' for _ in row)})", list(row.values()))
        # every line item, plus the agency totals (accounts in the pilot); no
        # other rollup and no memo line
        lines = {o["observation_id"] for o in r["observations"] if not (o["is_rollup"] or o["is_memo"])}
        extra = {row["canonical_account_id"] for row in rows if row["observation_id"] not in lines}
        self.assertTrue(lines <= {row["observation_id"] for row in rows})
        self.assertEqual(extra, {"ACC-NASA-TOTAL", "ACC-NSF-TOTAL"})

        res, h = self.query("NASA Science")
        got = {(o["fiscal_year"], o["stage"]): o["amount"] for o in h["observations"]}
        want = {(o["fiscal_year"], o["stage"]): o["amount"] for o in workbook_rows("Appropriations Observation")
                if o["canonical_account_id"] == "ACC-NASA-SCIENCE" and o["source_document_id"] == "SRC-CRPT-119HRPT652"}
        self.assertEqual(want, {(2026, "Enacted"): 7_250_000_000})
        self.assertEqual(got[(2026, "Enacted")], 7_250_000_000)          # dollars, not thousands
        self.assertEqual(h["observations"][0]["chamber"], "N/A")
        self.assertIn((2027, "House Reported"), got)                      # the column the pilot didn't load

    def test_an_account_row_without_an_account_is_refused(self):
        o = {"observation_id": "x", "account_match": "ambiguous", "canonical_account_id": None,
             "account_name_as_written": "Office of Inspector General"}
        with self.assertRaises(ValueError):
            S.observation_row(o, "SRC-CRPT-119HRPT652")


@unittest.skipUnless(openpyxl, "openpyxl not installed")
class LoadRefuses(unittest.TestCase):
    """Each mutation is something the loader must refuse rather than store."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def mutate(self, edit, data_only=True):
        wb = openpyxl.load_workbook(WORKBOOK, data_only=data_only)
        edit(wb)
        path = self.dir / "wb.xlsx"
        wb.save(path)
        return path

    def cell(self, wb, tab, row_id, col):
        ws = wb[tab]
        head = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == row_id:
                return row[head.index(col)]
        raise KeyError(row_id)

    def refuse(self, path, text):
        with self.assertRaises(S.LoadError) as cm:
            S.load(path, self.dir / "x.db")
        self.assertIn(text, str(cm.exception))
        self.assertFalse((self.dir / "x.db").exists())

    def test_unknown_enum_value(self):
        path = self.mutate(lambda wb: setattr(self.cell(wb, "Appropriations Observation", "OBS-0082",
                                                        "verification_status"), "value", "human_verified"))
        self.refuse(path, "CHECK constraint failed")

    def test_dangling_reference(self):
        path = self.mutate(lambda wb: setattr(self.cell(wb, "Appropriations Observation", "OBS-0082",
                                                        "source_document_id"), "value", "SRC-NOT-THERE"))
        self.refuse(path, "FOREIGN KEY constraint failed")

    def test_boolean_that_is_not_one(self):
        path = self.mutate(lambda wb: setattr(self.cell(wb, "Historical Name", "HN-0001", "human_reviewed"),
                                              "value", "yes"))
        self.refuse(path, "not a boolean")

    def test_display_list_drift(self):
        path = self.mutate(lambda wb: setattr(self.cell(wb, "Account", "ACC-NASA-EXPLORATION", "historical_names"),
                                              "value", "Deep Space Exploration Systems (FY2024 JES and earlier)"))
        self.refuse(path, "Account.historical_names")

    def test_formula_values_lost_on_save(self):
        # openpyxl (or any tool that doesn't recalculate) drops the cached
        # results of the bill_id/report_id lookup formulas on save
        path = self.mutate(lambda wb: None, data_only=False)
        self.refuse(path, "Bill Report Reference says")

    def test_failed_load_leaves_the_existing_database(self):
        db = self.dir / "x.db"
        S.load(WORKBOOK, db)
        before = db.read_bytes()
        bad = self.mutate(lambda wb: setattr(self.cell(wb, "Appropriations Observation", "OBS-0082",
                                                       "source_document_id"), "value", "SRC-NOT-THERE"))
        with self.assertRaises(S.LoadError):
            S.load(bad, db)
        self.assertEqual(db.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
