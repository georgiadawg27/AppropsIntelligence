"""
The workbook is built from data/staged.json by scripts/build_workbook.py (CI does it in the
data job). These tests build it and hold it to the reference build (tests/fixtures/workbook_values.json;
v38 plus the FY2024 House Labor-HHS draft), value for value,
and check that the build's own checks stop it on bad data.
"""

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import workbook_digest as W  # noqa: E402

SCRIPT = ROOT / "scripts" / "build_workbook.py"
STAGED = ROOT / "data" / "staged.json"
V38 = ROOT / "tests" / "fixtures" / "workbook_values.json"     # the reference build (v38 until the FY2024 House draft)
OUT = "build/Approps_Pilot_Schema_Loaded.xlsx"


def build(cwd):
    """Run the build script with cwd as the repo root it reads data/ from and writes build/ to."""
    (Path(cwd) / "build").mkdir(exist_ok=True)
    return subprocess.run([sys.executable, str(SCRIPT)], cwd=cwd, capture_output=True, text=True, timeout=300)


class BuiltWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        (Path(cls.tmp.name) / "data").mkdir()
        shutil.copy(STAGED, Path(cls.tmp.name) / "data" / "staged.json")
        r = build(cls.tmp.name)
        assert r.returncode == 0, r.stdout + r.stderr
        cls.path = Path(cls.tmp.name) / OUT
        cls.wb = openpyxl.load_workbook(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_matches_the_reference_workbook_sheet_by_sheet(self):
        want = json.loads(V38.read_text())["sheets"]
        self.assertEqual(W.compare(W.digest(self.path), want), [])
        self.assertEqual(list(want), ["Read Me", "Account", "Historical Name", "Confirmed Absence",
                                      "Appropriations Observation", "Account Relationship", "Source Document",
                                      "Bill Report Reference", "Validation Record", "Component"])

    def test_the_digest_sees_a_single_changed_value(self):
        with tempfile.TemporaryDirectory() as d:
            wb = openpyxl.load_workbook(self.path)
            wb["Appropriations Observation"]["H2"] = (wb["Appropriations Observation"]["H2"].value or 0) + 1
            wb.save(Path(d) / "x.xlsx")
            diffs = W.compare(W.digest(Path(d) / "x.xlsx"), json.loads(V38.read_text())["sheets"])
        self.assertEqual(len(diffs), 1)
        self.assertTrue(diffs[0].startswith("Appropriations Observation:"))

    def test_formulas_are_the_four_lookups_on_every_observation(self):
        n_obs = len(json.loads(STAGED.read_text())["observations"])
        formulas = [(ws.title, c.column_letter) for ws in self.wb.worksheets
                    for row in ws.iter_rows() for c in row if c.data_type == "f"]
        self.assertEqual(len(formulas), 4 * n_obs)
        self.assertEqual(len(formulas), 12_436)                # 3,109 observations x 4 lookups
        self.assertEqual({f for f in formulas}, {("Appropriations Observation", col) for col in "FGST"})

    def test_text_starting_with_equals_stays_text(self):
        cells = [c for ws in self.wb.worksheets for row in ws.iter_rows() for c in row
                 if isinstance(c.value, str) and c.value.startswith("=") and c.data_type != "f"]
        self.assertEqual(len(cells), 5)
        self.assertTrue(all(c.data_type == "s" and c.parent.title == "Validation Record" for c in cells))
        # and it is written to the file as a string, never as a formula (openpyxl names sheets sheet1..N in order)
        import zipfile
        n = self.wb.sheetnames.index("Validation Record") + 1
        with zipfile.ZipFile(self.path) as z:
            sheet = z.read(f"xl/worksheets/sheet{n}.xml").decode()
        self.assertIn("= minus the FY2022 advance", sheet)
        self.assertNotIn("<f>", sheet)


class BuildChecksStopTheBuild(unittest.TestCase):
    """Each kind of bad data stops the build (exit code != 0) and writes no workbook."""

    def run_with(self, change):
        data = json.loads(STAGED.read_text())
        change(data)
        with tempfile.TemporaryDirectory() as d:
            (Path(d) / "data").mkdir()
            (Path(d) / "data" / "staged.json").write_text(json.dumps(data))
            r = build(d)
            self.assertNotEqual(r.returncode, 0, "the build should have stopped")
            self.assertFalse((Path(d) / OUT).exists())
            return r.stdout + r.stderr

    def test_a_value_outside_its_allowed_list(self):
        def bad(d): d["observations"][0]["stage"] = "Conference"
        self.assertIn("values outside the allowed lists", self.run_with(bad))

    def test_a_broken_reference(self):
        def bad(d): d["observations"][0]["canonical_account_id"] = "ACC-NOPE"
        self.assertIn("unknown account", self.run_with(bad))

    def test_a_duplicate_id(self):
        def bad(d): d["observations"].append(dict(d["observations"][0]))
        self.assertIn("duplicate observation_id", self.run_with(bad))

    def test_an_effective_date(self):
        def bad(d): d["accounts"][0]["effective_start"] = 2017
        self.assertIn("effective_start is no longer a field", self.run_with(bad))

    def test_a_govinfo_landing_page(self):
        def bad(d): d["source_docs"][0]["url_or_identifier"] = "https://www.govinfo.gov/app/details/CRPT-119hrpt271"
        self.assertIn("govinfo landing page", self.run_with(bad))

    def test_a_status_the_rule_does_not_give(self):
        def bad(d):
            o = next(o for o in d["observations"] if o["verification_status"] == "flagged")
            o["verification_status"] = "auto-validated"
        self.assertIn("verification_status rule", self.run_with(bad))

    def test_a_resolution_left_blank(self):
        def bad(d):
            v = next(v for v in d["validations"] if v["human_review_status"] == "resolved")
            v["resolution"] = ""
        self.assertIn("resolved without a resolution", self.run_with(bad))

    def test_enactment_fields_on_a_row_that_is_not_enacted(self):
        def bad(d):
            r = next(r for r in d["bill_report_refs"] if r["stage"] != "Enacted")
            r["funding_type"] = "omnibus"
        self.assertIn("enactment fields are only for Enacted rows", self.run_with(bad))


if __name__ == "__main__":
    unittest.main()
