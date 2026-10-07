"""
Source links: what the site and the workbook show for a document.

Run:  python -m unittest tests.test_sources -v

- every source URL the site opens at a page ("#page=N") is a PDF, or on the
  short, documented exception list below;
- a document's public link (public_links.py) is derived at ingest, separate
  from the API fetch link -- Congress.gov for an Appropriations committee
  report found through its bill, else govinfo's content PDF, else a manual
  ingest's --source-url -- and kept only when its file has the stored sha256
  (Congress.gov answers come from responses recorded in
  tests/fixtures/congress_api/; no live calls);
- no stored or shown URL carries an api_key.
"""

import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import approps_store as S  # noqa: E402
import extract_approps as ex  # noqa: E402
import govinfo_ingest as g  # noqa: E402
import public_links as pl  # noqa: E402
from test_ingest import TempStore  # noqa: E402

WORKBOOK = S.reference_workbook()

# Sources cited at a page whose link is not a PDF. Each one is a known gap, not a pattern:
# the list must not grow without a reason written here.
PAGE_LINK_EXCEPTIONS = {
    # the Senate committee's download link for its FY2023 Labor-HHS draft explanatory statement
    # (S. 4659 was introduced, never reported, so govinfo has no report); a download link, which
    # browsers save rather than open at a page
    "SRC-EXPL-LHHS-FY2023-SENATE": "Senate committee download link (no govinfo report exists)",
    # CBO's web page for its H.R. 8845 cost estimate; the estimate's own PDF sits behind it
    "SRC-CBO-HR8845-FY2027": "CBO cost-estimate web page",
}


class Store(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.db = Path(cls.tmp.name) / "approps.db"
        S.load(WORKBOOK, cls.db)
        cls.conn = S.connect(cls.db, readonly=True)

    @classmethod
    def tearDownClass(cls):
        cls.conn.close()
        cls.tmp.cleanup()


class PageLinks(Store):
    def test_every_source_cited_at_a_page_links_to_a_pdf(self):
        rows = self.conn.execute(
            "SELECT DISTINCT d.document_id, d.url_or_identifier FROM source_document d "
            "JOIN appropriations_observation o ON o.source_document_id = d.document_id "
            "WHERE COALESCE(o.source_page, '') != ''").fetchall()
        self.assertGreater(len(rows), 20)
        not_pdf = {r[0]: r[1] for r in rows if not r[1].split("#")[0].lower().endswith(".pdf")}
        self.assertEqual(set(not_pdf), set(PAGE_LINK_EXCEPTIONS), not_pdf)

    def test_the_exceptions_are_still_what_they_say(self):
        # once one of them points at a PDF, it comes off the list
        for doc_id in PAGE_LINK_EXCEPTIONS:
            url = self.conn.execute("SELECT url_or_identifier FROM source_document WHERE document_id = ?",
                                    (doc_id,)).fetchone()[0]
            self.assertFalse(url.lower().endswith(".pdf"), doc_id)

    def test_the_nine_v35_sources_open_at_a_page(self):
        urls = dict(self.conn.execute("SELECT document_id, url_or_identifier FROM source_document"))
        for doc_id in ("SRC-CRPT-118SRPT84", "SRC-CRPT-118SRPT207", "SRC-CRPT-119SRPT55", "SRC-CRPT-118HRPT585",
                       "SRC-CRPT-119HRPT271", "SRC-CRPT-119HRPT696", "SRC-CRPT-117HRPT403",
                       "SRC-EXPL-LHHS-FY2026-ENACTED", "SRC-PLAW-119PUBL4"):
            self.assertRegex(urls[doc_id], r"^https://www\.govinfo\.gov/content/pkg/[^/]+/pdf/[^/]+\.pdf$", doc_id)


class NoApiKey(Store):
    def test_no_stored_url_carries_an_api_key(self):
        urls = [u for (u,) in self.conn.execute(
            "SELECT url_or_identifier FROM source_document UNION ALL "
            "SELECT bill_url FROM bill_report_reference UNION ALL SELECT report_jes_url FROM bill_report_reference")
            if u]
        self.assertGreater(len(urls), 30)
        self.assertEqual([u for u in urls if "api_key" in u.lower()], [])

    def test_the_site_shows_none(self):
        hits = [p.name for p in (ROOT / "docs").rglob("*") if p.is_file() and p.suffix in (".json", ".html", ".js")
                and "api_key" in p.read_text(errors="replace")]
        self.assertEqual(hits, [])

    def test_a_workbook_with_one_does_not_load(self):
        with tempfile.TemporaryDirectory() as d:
            wb = openpyxl.load_workbook(WORKBOOK, data_only=True)      # values: the lookup columns stay filled
            ws = wb["Source Document"]
            col = [c.value for c in ws[1]].index("url_or_identifier") + 1
            ws.cell(row=2, column=col, value="https://api.govinfo.gov/packages/X/pdf?api_key=abc123")
            bad = Path(d) / WORKBOOK.name
            wb.save(bad)
            with self.assertRaises(S.LoadError) as e:
                S.load(bad, Path(d) / "x.db")
            self.assertIn("carries an api_key", str(e.exception))
            self.assertNotIn("abc123", str(e.exception))


FIXTURES = json.loads((ROOT / "tests" / "fixtures" / "congress_api" / "responses.json").read_text())


def recorded(path, api_key):
    """public_links.congress_get, answered from responses recorded from the Congress.gov API."""
    assert api_key == "TESTKEY"
    return FIXTURES[path.lower()]


class Web:
    """public_links.http_get stand-in: these URLs serve these bytes; anything else is unreachable."""

    def __init__(self, files):
        self.files, self.asked = files, []

    def __call__(self, url, accept=None):
        self.asked.append(url)
        if url not in self.files:
            raise g.URLError("not served in this test")
        return self.files[url]


def pdf(tag):
    return b"%PDF-1.7 " + tag.encode()


def sha(b):
    return g.sha256_of(b)


def derive(package_id, entry, web, **kw):
    with mock.patch.object(pl, "congress_get", side_effect=recorded), mock.patch.object(pl, "http_get", web):
        return pl.derive(package_id, entry, api_key="TESTKEY", **kw)


class PublicLinks(unittest.TestCase):
    def test_a_report_found_through_its_bill(self):
        body = pdf("H.Rept. 119-271")
        url = "https://www.congress.gov/119/crpt/hrpt271/CRPT-119hrpt271.pdf"
        res = derive("CRPT-119hrpt271", {"hash": sha(body), "ingest_method": "govinfo_api"}, Web({url: body}),
                     bill_id="H.R.5304")
        self.assertEqual((res["public_url"], res["public_url_source"], res["link_needs_review"]), (url, "congress.gov", False))

    def test_a_report_under_a_pt1_granule_is_part_1(self):
        # govinfo files H.Rept. 118-585 under the granule CRPT-118hrpt585-pt1; Congress.gov's part 1 is CRPT-118hrpt585.pdf
        body = pdf("H.Rept. 118-585")
        url = "https://www.congress.gov/118/crpt/hrpt585/CRPT-118hrpt585.pdf"
        res = derive("CRPT-118hrpt585", {"hash": sha(body), "granule_id": "CRPT-118hrpt585-pt1"}, Web({url: body}),
                     bill_id="H.R.9029")
        self.assertEqual((res["public_url"], res["public_url_source"]), (url, "congress.gov"))

    def test_a_two_part_report_is_matched_by_part(self):
        # H.Rept. 118-74 (recorded): part 1 Science, Space, and Technology; part 2 Agriculture
        text = FIXTURES["committee-report/118/hrpt/74/text"]
        self.assertEqual(pl.pdf_for_part(text, "CRPT-118hrpt74", 1), "https://www.congress.gov/118/crpt/hrpt74/CRPT-118hrpt74.pdf")
        self.assertEqual(pl.pdf_for_part(text, "CRPT-118hrpt74", 2), "https://www.congress.gov/118/crpt/hrpt74/CRPT-118hrpt74-pt2.pdf")
        self.assertIsNone(pl.pdf_for_part(text, "CRPT-118hrpt74", 3))
        self.assertEqual(pl.part_of("CRPT-118hrpt74-pt2"), 2)
        # neither part is an Appropriations Committee report: no Congress.gov link; govinfo's part-2 granule PDF instead
        body = pdf("H.Rept. 118-74 part 2")
        gi = "https://www.govinfo.gov/content/pkg/CRPT-118hrpt74/pdf/CRPT-118hrpt74-pt2.pdf"
        res = derive("CRPT-118hrpt74", {"hash": sha(body), "granule_id": "CRPT-118hrpt74-pt2"}, Web({gi: body}), bill_id="H.R.1713")
        self.assertEqual((res["public_url"], res["public_url_source"]), (gi, "govinfo"))
        self.assertEqual(res["tried"][0]["source"], "congress.gov")
        self.assertIsNone(res["tried"][0]["url"])

    def test_a_bill_with_a_non_appropriations_report(self):
        # H.R. 2882 (118th) was the Udall Foundation Reauthorization Act before it carried FY2024
        # appropriations: its one report, H.Rept. 118-364, is Natural Resources'
        with mock.patch.object(pl, "congress_get", side_effect=recorded):
            url, note = pl.congress_report_pdf((118, "hr", 2882), "TESTKEY")
        self.assertIsNone(url)
        self.assertEqual(note, "no Appropriations Committee report for the bill")

    def test_a_record_explanatory_statement_via_govinfo(self):
        # the FY2026 Labor-HHS explanatory statement, printed in the Congressional Record, ingested by hand
        # with its govinfo landing page: its public link is the article's content PDF
        body = pdf("CREC-2026-01-22-pt2-PgH1353-2")
        gi = "https://www.govinfo.gov/content/pkg/CREC-2026-01-22/pdf/CREC-2026-01-22-pt2-PgH1353-2.pdf"
        entry = {"hash": sha(body), "ingest_method": "manual", "doc_type": "jes",
                 "source_url": "https://www.govinfo.gov/app/details/CREC-2026-01-22/CREC-2026-01-22-pt2-PgH1353-2"}
        res = derive("MANUAL-LHHS-FY2026-Enacted-jes-e44f7662", entry, Web({gi: body}))
        self.assertEqual((res["public_url"], res["public_url_source"]), (gi, "govinfo"))
        # a public law: govinfo's package PDF
        body = pdf("P.L. 119-4")
        gi = "https://www.govinfo.gov/content/pkg/PLAW-119publ4/pdf/PLAW-119publ4.pdf"
        res = derive("PLAW-119publ4", {"hash": sha(body), "ingest_method": "govinfo_api"}, Web({gi: body}))
        self.assertEqual((res["public_url"], res["public_url_source"]), (gi, "govinfo"))

    def test_a_committee_site_document_keeps_its_source_url(self):
        body = pdf("FY2023 Senate LHHS draft")
        url = "https://www.appropriations.senate.gov/imo/media/doc/lhhsfy23rept.pdf"
        res = derive("MANUAL-LHHS-FY2023-SenateReported-explanatory_statement-88301550",
                     {"hash": sha(body), "ingest_method": "manual", "source_url": url}, Web({url: body}))
        self.assertEqual((res["public_url"], res["public_url_source"]), (url, "manual"))

    def test_a_sha_mismatch_goes_to_needs_review(self):
        url = "https://www.congress.gov/119/crpt/hrpt271/CRPT-119hrpt271.pdf"
        gi = "https://www.govinfo.gov/content/pkg/CRPT-119hrpt271/pdf/CRPT-119hrpt271.pdf"
        entry = {"hash": sha(pdf("what we ingested")), "ingest_method": "govinfo_api"}
        res = derive("CRPT-119hrpt271", entry, Web({url: pdf("a revised print"), gi: b"<html>not a pdf</html>"}),
                     bill_id="H.R.5304")
        self.assertEqual((res["public_url"], res["link_needs_review"]), (None, True))
        self.assertEqual([(t["source"], t["matched"]) for t in res["tried"]], [("congress.gov", False), ("govinfo", False)])
        pl.apply(entry, res)
        self.assertEqual(pl.needing_review({"CRPT-119hrpt271": entry, "X": {"link_needs_review": False}}), ["CRPT-119hrpt271"])
        sd = ex.source_document_fields("CRPT-119hrpt271", "sha", {"document_type": "committee_report", "stage": "House Reported",
                                       "congress_session": "119", "report_id": None, "bill_id": None}, entry, [])
        self.assertEqual((sd["url_or_identifier"], sd["link_needs_review"]), ("govinfo:CRPT-119hrpt271", True))

    def test_a_manual_ingest_without_its_source_url_is_refused(self):
        with tempfile.TemporaryDirectory() as d:
            f = Path(d) / "es.pdf"
            f.write_bytes(pdf("manual"))
            kw = dict(subcommittee="LHHS", fiscal_year=2023, stage="Senate Reported", doc_type="explanatory_statement",
                      advance_copy=False)
            with mock.patch.object(g, "STORE_DIR", Path(d) / "store"), \
                    mock.patch.object(g, "MANIFEST_PATH", Path(d) / "store" / "manifest.json"), \
                    mock.patch.object(pl, "http_get", Web({})):
                with self.assertRaises(ValueError):
                    g.ingest_local(f, {}, **kw)
                with self.assertRaises(SystemExit), mock.patch("sys.stderr"):
                    g.main_ingest_local([str(f), "--subcommittee", "LHHS", "--fiscal-year", "2023", "--stage",
                                         "Senate Reported", "--doc-type", "explanatory_statement"])
                m = {}
                res = g.ingest_local(f, m, public=False, **kw)              # published nowhere: no link, nothing to review
                self.assertEqual((m[res["package_id"]]["public_url"], m[res["package_id"]]["link_needs_review"]), (None, False))
                self.assertEqual(m[res["package_id"]]["ingest_method"], "manual")
                with self.assertRaises(ValueError):
                    g.ingest_local(f, {}, source_url="https://api.govinfo.gov/packages/X/pdf?api_key=abc", **kw)

    def test_no_stored_url_carries_an_api_key(self):
        # a full fetch: the API fetch carries the key; the manifest stores the fetch link without it and the public link
        body = pdf("H.Rept. 119-271")
        public = "https://www.congress.gov/119/crpt/hrpt271/CRPT-119hrpt271.pdf"
        asked = []

        class Resp:
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def read(self):
                return body

        def urlopen(req):
            asked.append(req.full_url)
            return Resp()
        with TempStore(http_get=Web({public: body})) as store, \
                mock.patch.object(g, "urlopen", side_effect=urlopen), \
                mock.patch.object(g, "api_get", return_value={"collectionCode": "CRPT", "download": {}}), \
                mock.patch.object(pl, "congress_get", side_effect=recorded):
            manifest = {}
            g.fetch_and_store("CRPT-119hrpt271", "TESTKEY", manifest, bill_id="BILLS-119hr5304rh")
            g.save_manifest(manifest)
            text = (store / "manifest.json").read_text()
        self.assertTrue(all("api_key=TESTKEY" in u for u in asked))
        e = manifest["CRPT-119hrpt271"]
        self.assertEqual((e["public_url"], e["public_url_source"]), (public, "congress.gov"))
        self.assertEqual(e["fetch_link"], f"{g.API_BASE}/packages/CRPT-119hrpt271/pdf")
        self.assertNotIn("api_key", text)
        self.assertNotIn("TESTKEY", text)
        self.assertNotIn("api_key", json.dumps(FIXTURES))


if __name__ == "__main__":
    unittest.main()
