"""
v33: the reference workbook with the Labor-HHS account list, and what came with it --
parent accounts, Source Document / Bill Report Reference notes, every Data
Dictionary document type, superseded_by_observation_id, the reconciliation
rule, headline lines that must sit in their headline's document, and the one
verification_status rule.

Run:  python -m unittest tests.test_v33 -v

Needs openpyxl (loading only).
"""

import contextlib
import io
import shutil
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import approps_store as S  # noqa: E402
import govinfo_ingest as g  # noqa: E402
import validate_approps as V  # noqa: E402

try:
    import openpyxl
except ImportError:                                  # pragma: no cover
    openpyxl = None

WORKBOOK = S.reference_workbook()


@unittest.skipUnless(openpyxl, "openpyxl not installed")
class V33Store(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.db = Path(cls.tmp.name) / "approps.db"
        cls.report = S.load(WORKBOOK, cls.db)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def setUp(self):
        self.conn = S.connect(self.db)

    def tearDown(self):
        self.conn.close()

    def copy(self):
        path = Path(self.tmp.name) / f"{self.id().rsplit('.', 1)[-1]}.db"
        shutil.copy(self.db, path)
        return S.connect(path)

    def test_counts(self):
        # the workbook is built from data/staged.json (v38), not read from a committed file
        self.assertEqual(WORKBOOK, S.BUILT_WORKBOOK)
        self.assertEqual(S.data_version(WORKBOOK), "v38")
        # FY2024 House Labor-HHS draft (SRC-EXPL-LHHS-FY2024-HOUSE): + 1 document, 1 reference, 101 observations, 2 absences, 473 records; then H.R. 5894's bill text: + 1 document, 1 observation (NEF rescission), 3 records; then law_text: + 374 records (328 pass, 46 info); then the FY2022 backfill (+ 2 documents, 4 references, 378 observations) and FY2023 Medicaid's derived headline and views (+ 12 observations, - 4 absences retired)
        self.assertEqual(self.report["rows"], {
            "account": 130, "historical_name": 14, "source_document": 39, "bill_report_reference": 63,
            "component": 22, "appropriations_observation": 3131, "confirmed_absence": 230,
            "account_relationship": 11, "validation_record": 9496})
        self.assertEqual(self.report["warnings"], [])

    def test_formula_looking_text_is_read_as_text(self):
        # five expected_result values start with "=": stored as text in the workbook, read as text
        rows = dict(self.conn.execute("SELECT validation_id, expected_result FROM validation_record "
                                      "WHERE expected_result LIKE '=%'").fetchall())
        self.assertEqual(sorted(rows), ["VAL-LHHS-02338", "VAL-LHHS-02352", "VAL-LHHS-02362",
                                        "VAL-LHHS-02378", "VAL-LHHS-02388"])
        self.assertTrue(rows["VAL-LHHS-02338"].startswith("= minus the FY2022 advance"))

    def test_notes_columns(self):
        brr = self.conn.execute("SELECT notes FROM bill_report_reference WHERE reference_id = "
                                "'BR-LHHS-FY2023-SENATE'").fetchone()[0]
        self.assertIn("never reported", brr)
        doc = self.conn.execute("SELECT notes FROM source_document WHERE document_id = "
                                "'SRC-EXPL-LHHS-FY2023-SENATE'").fetchone()[0]
        self.assertIn("883015501e56e3437477e750dcaba220bf13468eba2a3a8f736f628be7152e18", doc)

    def test_stage_note_rides_with_the_history_and_the_grid(self):
        key = "2023|Senate Reported"
        h = S.history(self.conn, "ACC-HHS-NIH-TOTAL")
        self.assertIn("never reported", h["stage_notes"][key])
        self.assertIn(key, S.subcommittee_grid(self.conn, "LHHS")["stage_notes"])
        # CJS has no notes: nothing new in its history or grid (its export stays as it was)
        self.assertNotIn("stage_notes", S.history(self.conn, "ACC-NASA-SCIENCE"))
        self.assertNotIn("stage_notes", S.subcommittee_grid(self.conn, "CJS"))
        o = next(o for o in h["observations"] if o["source_document_id"] == "SRC-EXPL-LHHS-FY2023-SENATE")
        self.assertIn("committee draft", o["document_notes"])

    # ---- v34: statuses and relationships ------------------------------------------------

    FLAGGED_FOR_REVIEW = ["OBS-LHHS-0434", "OBS-LHHS-0440", "OBS-LHHS-0457", "OBS-LHHS-0465", "OBS-LHHS-0466",
                          "OBS-LHHS-0467", "OBS-LHHS-0468", "OBS-LHHS-0471", "OBS-LHHS-0472", "OBS-LHHS-0477",
                          "OBS-LHHS-0478", "OBS-LHHS-0484", "OBS-LHHS-0485", "OBS-LHHS-0488"]

    def checks(self, oid):
        return [tuple(r) for r in self.conn.execute(
            "SELECT rule_applied, result, expected_result, human_review_status, resolution FROM validation_record "
            "WHERE observation_id = ?", (oid,))]

    def test_a_resolved_check_no_longer_counts(self):
        fail = ("structural", "fail", "1 = sum of [a]")
        ok = ("table_total", "pass", "2 = sum of [b]")
        self.assertEqual(V.verification_status(0.95, [fail, ok]), "flagged")
        self.assertEqual(V.verification_status(0.95, [fail + ("resolved", "parse artifact"), ok]), "auto-validated")
        self.assertEqual(V.verification_status(0.95, [fail + ("resolved", " "), ok]), "flagged")     # no resolution written
        self.assertEqual(V.verification_status(0.95, [fail + ("pending", "x"), ok]), "flagged")
        self.assertEqual(V.verification_status(0.50, [fail + ("resolved", "parse artifact"), ok]), "unverified")
        # a resolved record neither flags nor confirms
        xd = ("cross_document", "flag", "another document prints 3")
        self.assertEqual(V.verification_status(0.95, [xd + ("resolved", "explained"), ("source_text", "pass")]), "unverified")
        self.assertEqual(V.verification_status(0.95, [("cross_document", "pass") + (None, "resolved", "x")]), "unverified")

    def test_law_text_confirms_and_info_counts_neither_way(self):
        self.assertIn("law_text", V.CONFIRMING_RULES)
        src = ("source_text", "pass")
        self.assertEqual(V.verification_status(0.95, [src, ("law_text", "pass", "H.R. 2617 div. H ...")]), "auto-validated")
        self.assertEqual(V.verification_status(0.95, [src, ("law_text", "info", "H.R. 2617 div. H ...")]), "unverified")
        self.assertEqual(V.verification_status(0.95, [("table_total", "pass"), ("law_text", "info")]), "auto-validated")
        self.assertEqual(V.verification_status(0.50, [src, ("law_text", "pass")]), "auto-validated")  # another document confirms below 0.90
        rows = self.conn.execute("SELECT result, count(*) FROM validation_record WHERE rule_applied = 'law_text' "
                                 "GROUP BY result").fetchall()
        # + LIHEAP FY2023, matched across divisions H + N; + FY2022 (H.R. 2471 div. H) and FY2023 Medicaid's derived headline;
        # + Refugee FY2022's other-law note (an info law_text record from the table, not the law); PHSSEF FY2022 matches the sum
        # of its heading's four paragraphs; earlier records the paragraph sum now matches updated to pass (owner, 2026-10-08)
        self.assertEqual(dict(rows), {"pass": 404, "info": 47})
        # never pending: an info is a recorded difference, not a question
        self.assertEqual(self.conn.execute("SELECT count(*) FROM validation_record WHERE rule_applied = 'law_text' "
                                           "AND human_review_status IS NOT NULL AND human_review_status <> ''").fetchone()[0], 0)

    def test_another_document_confirms_below_the_threshold(self):
        # below 0.90 a passing check against another document confirms; the stored confidence is unchanged
        src, total = ("source_text", "pass"), ("table_total", "pass")
        self.assertEqual(V.verification_status(0.5, [src, total]), "unverified")
        self.assertEqual(V.verification_status(0.5, [src, total, ("cross_document", "pass")]), "auto-validated")
        self.assertEqual(V.verification_status(0.5, [src, ("law_text", "pass")]), "auto-validated")
        self.assertEqual(V.verification_status(0.5, [src, ("law_text", "info")]), "unverified")
        self.assertEqual(V.verification_status(0.5, [("law_text", "pass"), ("structural", "fail")]), "flagged")
        # every cross_document pass that lifts a figure below 0.90 cites a document other than the figure's own
        import re
        for oid, src_doc, exp in self.conn.execute(
                "SELECT o.observation_id, o.source_document_id, v.expected_result FROM appropriations_observation o "
                "JOIN validation_record v USING (observation_id) WHERE o.confidence < 0.9 "
                "AND v.rule_applied = 'cross_document' AND v.result = 'pass'").fetchall():
            m = re.search(r"(\d{3})[HS]RPT(\d+)", src_doc)
            self.assertNotIn(src_doc, exp, oid)
            if m:
                self.assertNotIn(f"{m.group(1)}-{m.group(2)} p", exp, oid)
        n = self.conn.execute("SELECT count(*) FROM appropriations_observation WHERE confidence < 0.9 "
                              "AND verification_status = 'auto-validated'").fetchone()[0]
        self.assertEqual(n, 115)                        # + 45 FY2022 figures another document confirms

    def test_the_status_rule_reproduces_every_status(self):
        n = 0
        for oid, conf, status in self.conn.execute(
                "SELECT observation_id, confidence, verification_status FROM appropriations_observation "
                "WHERE verification_status NOT IN ('human-verified', 'provisional', 'superseded')").fetchall():
            self.assertEqual(V.verification_status(conf, self.checks(oid)), status, oid)
            n += 1
        self.assertEqual(n, 2257)                       # + the 101 FY2024 House draft rows and the NEF rescission, by the same rule; + FY2022 and the FY2023/FY2024 Medicaid lines

    def test_the_fourteen_deliberate_flags(self):
        # the two Title II scope totals stay flagged (pending); the owner resolved the twelve FY2025 Enacted
        # estimates (PR #29), which leaves them unverified -- nothing else confirms an estimate
        got = dict(self.conn.execute("SELECT observation_id, verification_status FROM appropriations_observation "
                                     "WHERE observation_id IN (%s)" % ",".join("?" * 14), self.FLAGGED_FOR_REVIEW).fetchall())
        self.assertEqual(got, {oid: "flagged" if oid in ("OBS-LHHS-0434", "OBS-LHHS-0440") else "unverified"
                               for oid in self.FLAGGED_FOR_REVIEW})
        for oid in self.FLAGGED_FOR_REVIEW:
            # none has a failed check: each was flagged by a review flag (a scope question or a disagreeing document)
            self.assertNotIn("fail", [c[1] for c in self.checks(oid)])
            self.assertTrue(any(V.review_flag(*c[:3]) for c in self.checks(oid)), oid)
            self.assertEqual(V.verification_status(0.95, [c[:3] for c in self.checks(oid)]), "flagged", oid)
            self.assertEqual(V.verification_status(0.95, self.checks(oid)), got[oid], oid)

    def test_a_routine_semantic_flag_is_not_a_review_flag(self):
        self.assertFalse(V.review_flag("semantic", "flag", "amount_type fits the row label (advance)"))
        self.assertTrue(V.review_flag("semantic", "flag", "a dollar level set in law; P.L. 119-4 sec. 1101"))
        self.assertTrue(V.review_flag("cross_document", "flag"))
        self.assertFalse(V.review_flag("structural", "flag", "anything"))
        self.assertEqual(V.verification_status(0.95, [("table_total", "pass"), ("semantic", "flag",
                                                       "amount_type fits the row label (advance)")]), "unverified")

    def test_the_cj_is_a_source_document(self):
        self.assertEqual(self.conn.execute("SELECT document_type, stage FROM source_document WHERE document_id = "
                                           "'SRC-CJ-AHA-FY2026'").fetchone()[:], ("congressional_budget_justification",
                                                                                  "President's Budget"))

    def test_a_relationship_shows_on_both_accounts_with_its_cj_pages(self):
        from_side = {r["relationship_id"]: r for r in S.history(self.conn, "ACC-HHS-NIH-NIEHS")["relationships"]}
        to_side = {r["relationship_id"]: r for r in S.history(self.conn, "ACC-HHS-AHA-TOTAL")["relationships"]}
        r = from_side["REL-LHHS-0011"]
        self.assertEqual((r["direction"], r["relationship_type"], r["other_account_id"]),
                         ("from", "moved_reclassified", "ACC-HHS-AHA-TOTAL"))
        self.assertEqual(r["cites"], [{"document_id": "SRC-CJ-AHA-FY2026",
                                       "url": "https://www.hhs.gov/sites/default/files/fy-2026-aha-cj.pdf",
                                       "pages": ["p.11", "p.13"]}])
        self.assertEqual((to_side["REL-LHHS-0011"]["direction"], to_side["REL-LHHS-0011"]["other_account_id"]),
                         ("to", "ACC-HHS-NIH-NIEHS"))
        self.assertEqual(set(to_side), {"REL-LHHS-0001", "REL-LHHS-0002", "REL-LHHS-0011", "REL-LHHS-0012",
                                        "REL-LHHS-0013", "REL-LHHS-0014", "REL-LHHS-0015"})   # 0003 / 0004 retired

    # ---- parent_account_id -----------------------------------------------------------

    def test_hrsa_total_does_not_count_health_centers_twice(self):
        members = S.agency_members(self.conn, "ACC-HHS-HRSA-TOTAL")
        self.assertIn("ACC-HHS-HRSA-PRIMARY-CARE", members)
        self.assertNotIn("ACC-HHS-HRSA-HEALTH-CENTERS", members)
        for fy, stage in ((2023, "Enacted"), (2024, "Senate Reported"), (2025, "Enacted")):
            total, have, lack = S.agency_sum(self.conn, "ACC-HHS-HRSA-TOTAL", fy, stage)
            self.assertEqual(lack, [], (fy, stage))
            # the members add up to the printed HRSA total exactly; with Health Centers added again they wouldn't
            self.assertEqual(total, S.cell_total(self.conn, "ACC-HHS-HRSA-TOTAL", fy, stage), (fy, stage))
            self.assertGreater(S.cell_total(self.conn, "ACC-HHS-HRSA-HEALTH-CENTERS", fy, stage), 0)

    def test_a_child_sits_under_its_parent_in_the_grid(self):
        grid = S.subcommittee_grid(self.conn, "LHHS")
        ids = [r["account"]["canonical_account_id"] for r in grid["rows"]]
        for parent, child in (("ACC-HHS-HRSA-PRIMARY-CARE", "ACC-HHS-HRSA-HEALTH-CENTERS"),
                              ("ACC-HHS-ACF-CFSP", "ACC-HHS-ACF-HEAD-START")):
            self.assertEqual(ids.index(child), ids.index(parent) + 1)
            row = grid["rows"][ids.index(parent)]
            self.assertEqual(row["children"], [child])
            self.assertEqual(grid["rows"][ids.index(child)]["account"]["parent_account_id"], parent)
        # an account with no parent carries no parent key (CJS's export is unchanged)
        self.assertNotIn("parent_account_id", S.subcommittee_grid(self.conn, "CJS")["rows"][0]["account"])

    def test_title_totals_leave_children_out(self):
        import importlib.util
        spec = importlib.util.spec_from_file_location("tt", ROOT / "reference" / "review" / "title_totals.py")
        tt = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(tt)
        base, _, _ = tt.cell_lines(self.conn, "Title II", 2024, "Senate Reported", subcommittee="LHHS")
        self.assertNotIn("ACC-HHS-HRSA-HEALTH-CENTERS", {r["canonical_account_id"] for r in base})
        self.assertIn("ACC-HHS-HRSA-PRIMARY-CARE", {r["canonical_account_id"] for r in base})

    def test_nih_fy2024_senate_keeps_cures_as_its_own_row(self):
        self.assertEqual(S.cell_total(self.conn, "ACC-HHS-NIH-TOTAL", 2024, "Senate Reported"), 47_811_518_000)
        self.assertIsNotNone(S.cell_total(self.conn, "ACC-HHS-NIH-CURES", 2024, "Senate Reported"))

    # ---- the Data Dictionary in full ---------------------------------------------------

    def test_every_dictionary_document_type_is_accepted(self):
        conn = self.copy()
        for i, t in enumerate(("congressional_budget_justification", "crs_report", "omb_public_budget_database")):
            conn.execute("INSERT INTO source_document (document_id, source_agency, url_or_identifier, document_type, "
                         "fiscal_year, stage) VALUES (?, 'HHS', 'x', ?, 2026, 'President''s Budget')", (f"SRC-T{i}", t))
        with self.assertRaises(sqlite3.IntegrityError):
            conn.execute("INSERT INTO source_document (document_id, source_agency, url_or_identifier, document_type, "
                         "fiscal_year, stage) VALUES ('SRC-BAD', 'HHS', 'x', 'blog_post', 2026, 'Enacted')")
        conn.close()

    def test_the_reconciliation_rule_loads(self):
        conn = self.copy()
        conn.execute("INSERT INTO validation_record (validation_id, observation_id, rule_applied, result) "
                     "VALUES ('VAL-T', 'OBS-0082', 'advance_copy_reconciliation', 'flag')")
        conn.close()

    def test_superseded_by(self):
        conn = self.copy()
        old = dict(conn.execute("SELECT * FROM appropriations_observation WHERE observation_id = 'OBS-0082'").fetchone())
        # a superseded row must name what replaced it, and only a superseded row may
        with self.assertRaises(sqlite3.IntegrityError):
            conn.execute("UPDATE appropriations_observation SET verification_status = 'superseded' "
                         "WHERE observation_id = 'OBS-0082'")
        with self.assertRaises(sqlite3.IntegrityError):
            conn.execute("UPDATE appropriations_observation SET superseded_by_observation_id = 'OBS-0083' "
                         "WHERE observation_id = 'OBS-0082'")
        # the official value replaces it: both rows stay, the same fact twice, one current
        with conn:
            conn.execute("UPDATE appropriations_observation SET verification_status = 'superseded', "
                         "superseded_by_observation_id = 'OBS-NEW' WHERE observation_id = 'OBS-0082'")
            new = dict(old, observation_id="OBS-NEW", amount=old["amount"] + 1000, superseded_by_observation_id=None)
            conn.execute(f"INSERT INTO appropriations_observation ({', '.join(new)}) VALUES "
                         f"({', '.join('?' for _ in new)})", list(new.values()))
        self.assertEqual(S.superseded_errors(conn), [])
        h = S.history(conn, old["canonical_account_id"])
        self.assertEqual({o["observation_id"] for o in h["observations"]} & {"OBS-0082", "OBS-NEW"},
                         {"OBS-0082", "OBS-NEW"})
        cell = next(r for r in S.history_grid(h)["rows"] if r["fiscal_year"] == old["fiscal_year"])["cells"][old["stage"]]
        shown = [o["observation_id"] for line in cell for o in line["observations"]]
        self.assertIn("OBS-NEW", shown)
        self.assertNotIn("OBS-0082", shown)
        self.assertEqual(S.cell_total(conn, old["canonical_account_id"], old["fiscal_year"], old["stage"],
                                      old["amount_type"]), new["amount"])
        conn.close()


@unittest.skipUnless(openpyxl, "openpyxl not installed")
class V33LoadRefuses(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_a_headline_in_another_document_is_a_load_error(self):
        wb = openpyxl.load_workbook(WORKBOOK, data_only=True)      # the lookup columns' values, as the loader reads them
        ws = wb["Appropriations Observation"]
        head = [c.value for c in ws[1]]
        col = {h: i for i, h in enumerate(head)}
        rows = [[c for c in r] for r in ws.iter_rows(min_row=2)]
        view = next(r for r in rows if r[col["headline_observation_id"]].value)
        doc = view[col["source_document_id"]].value
        other = next(r for r in rows if r[col["component"]].value is None
                     and r[col["source_document_id"]].value != doc
                     and r[col["canonical_account_id"]].value == view[col["canonical_account_id"]].value)
        view[col["headline_observation_id"]].value = other[col["observation_id"]].value
        path = self.dir / "wb.xlsx"
        wb.save(path)
        with self.assertRaises(S.LoadError) as cm, contextlib.redirect_stdout(io.StringIO()):
            S.load(path, self.dir / "x.db")
        self.assertIn("source_document_id", str(cm.exception))
        self.assertIn(view[col["observation_id"]].value, str(cm.exception))
        self.assertFalse((self.dir / "x.db").exists())


class VerificationStatusRule(unittest.TestCase):
    ok = [("source_text", "pass"), ("unit", "pass"), ("table_total", "pass")]

    def test_auto_validated_needs_confidence_and_every_check(self):
        self.assertEqual(V.verification_status(0.95, self.ok), "auto-validated")
        self.assertEqual(V.verification_status(0.90, self.ok), "auto-validated")
        self.assertEqual(V.verification_status(0.50, self.ok), "unverified")       # a pass doesn't lift a low confidence
        self.assertEqual(V.verification_status(0.95, self.ok + [("semantic", "flag")]), "unverified")
        self.assertEqual(V.verification_status(0.95, self.ok + [("structural", "fail")]), "flagged")

    def test_something_has_to_confirm_the_figure(self):
        self.assertEqual(V.verification_status(0.95, [("source_text", "pass"), ("unit", "pass")]), "unverified")
        self.assertEqual(V.verification_status(0.95, []), "unverified")
        self.assertEqual(V.verification_status(0.95, [("cross_document", "pass")]), "auto-validated")


class DownloadCleansTheKey(unittest.TestCase):
    def test_download_uses_the_cleaned_key(self):
        seen = []

        class Resp:
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def read(self):
                return b"%PDF"

        with mock.patch.object(g, "urlopen", side_effect=lambda req: seen.append(req.full_url) or Resp()):
            g.download("https://api.govinfo.gov/packages/X/pdf", ' "abc123" \n')
        self.assertEqual(seen, ["https://api.govinfo.gov/packages/X/pdf?api_key=abc123"])



class BudgetJustificationIngest(unittest.TestCase):
    def test_a_cj_ingests_and_is_never_an_advance_copy(self):
        with tempfile.TemporaryDirectory() as d:
            pdf = Path(d) / "cj.pdf"
            pdf.write_bytes(b"%PDF-1.7 test")
            with mock.patch.object(g, "STORE_DIR", Path(d) / "store"), \
                    mock.patch.object(g.public_links, "http_get", side_effect=g.URLError("offline test")):
                m = {}
                res = g.ingest_local(pdf, m, subcommittee="LHHS", fiscal_year=2026, stage="President's Budget",
                                     doc_type="congressional_budget_justification", advance_copy=False,
                                     source_url="https://www.hhs.gov/sites/default/files/fy-2026-aha-cj.pdf")
                self.assertEqual(m[res["package_id"]]["confirmation_status"], "no_official_counterpart")
                with self.assertRaises(ValueError):
                    g.ingest_local(pdf, {}, subcommittee="LHHS", fiscal_year=2026, stage="President's Budget",
                                   doc_type="congressional_budget_justification", advance_copy=True,
                                   source_url="https://www.hhs.gov/sites/default/files/fy-2026-aha-cj.pdf")
        import extract_approps as ex
        self.assertEqual(ex.describe_package(res["package_id"], m[res["package_id"]])["document_type"],
                         "congressional_budget_justification")


if __name__ == "__main__":
    unittest.main()
