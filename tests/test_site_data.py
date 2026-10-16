"""
The site reads data/staged.json: the workbook is built from it (approps_store.reference_workbook),
the export publishes that build as docs/Approps_Pilot_Schema_Loaded.xlsx, and nothing reads a
committed workbook. These tests hold the switch to what the site showed before it.
"""

import collections
import json
import re
import tempfile
import unittest
from pathlib import Path

import approps_store as S
import export_static as E

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"

# grid-state counts and flagged cells, measured on the export from the committed v37 workbook
# just before the switch (the grid's headline cell is flagged when its observation's
# verification_status is "flagged"); then the FY2024 House Labor-HHS draft: its 100 cells moved
# from not yet collected to figures and states, 4 of them flagged; then H.R. 5894: the NEF rescission moved from
# missing to a figure, and CDC Global Health became flagged (its bill-text disagreement); then the owner's
# triage (PR #29): resolved records no longer count, 78 flagged cells -> 5 (71 unverified, 2 auto-validated);
# then the FY2022 Labor-HHS backfill (+400 cells) and 'no figure for this year' (a proposed account before its first figure)
# the published grid's counts, regenerated with the data (scripts/grid_counts.py): a change that moves them shows in
# the fixture's diff
GRID_COUNTS = json.loads((ROOT / "tests" / "fixtures" / "grid_counts.json").read_text())
STATE_COUNTS = GRID_COUNTS["state_counts"]
FLAGGED_CELLS = GRID_COUNTS["flagged_cells"]


def grid(sc):
    return E.read_grid(DOCS / "data", sc)


def headline_observations(g):
    """The observation each grid cell shows (the page's headOf: the budget authority line with
    no component, else the first line), over every row, title total and bill total."""
    rows = g["rows"] + [t["total"] for t in g.get("titles", []) if t.get("total")] + \
        ([g["bill_total"]] if g.get("bill_total") else [])
    for row in rows:
        for key, cell in row["cells"].items():
            if not cell["lines"]:
                continue
            h = next((l for l in cell["lines"] if l["amount_type"] == "budget authority" and not l["component"]),
                     cell["lines"][0])
            if h["observations"]:
                yield row, key, h["observations"][0]


class SiteReadsData(unittest.TestCase):
    def test_no_committed_workbook(self):
        self.assertEqual(list((ROOT / "reference").glob("*.xlsx")), [])
        self.assertEqual(S.reference_workbook(), S.BUILT_WORKBOOK)

    def test_index_names_the_data(self):
        src = json.loads((DOCS / "data" / "index.json").read_text())["source"]
        self.assertEqual(src["data"], "data/staged.json")
        self.assertEqual(src["version"], "v38")
        self.assertEqual(src["download"], "Approps_Pilot_Schema_Loaded.xlsx")
        import hashlib
        self.assertEqual(src["data_sha256"], hashlib.sha256(S.STAGED.read_bytes()).hexdigest())

    def test_grid_state_counts_unchanged(self):
        for sc, want in STATE_COUNTS.items():
            self.assertEqual(grid(sc)["state_counts"], want, sc)

    def test_flagged_cell_count_unchanged(self):
        for sc, want in FLAGGED_CELLS.items():
            got = sum(1 for _, _, o in headline_observations(grid(sc)) if o["verification_status"] == "flagged")
            self.assertEqual(got, want, sc)

    def test_review_status_comes_with_every_cell(self):
        # Reviewer mode reads human_review_status from each observation's validation records. The owner's
        # triage (PR #29) resolved 452 of the 462 pending records; 10 stay pending. The FY2022 backfill adds 27 pending
        # (sum and cross-document questions) and 90 resolved at creation by the triage's standard reasons; the owner's
        # rules 3 and 5 (2026-10-08) resolve 13 more (short sums explained exactly, the request baseline) and 16 sums
        # pass (the approved former names, the derived CDC-Wide headline); 11 stay pending. The bill and law text
        # check's mismatches and ambiguous results (VAL-TXT-) wait for review too, and so do measure_reconcile's
        # failed bases and printed views that differ from the computed ones (VAL-MR-; one VAL-MRC- item per cause of the printed views that differ)
        data = json.loads(S.STAGED.read_text())
        # the backfill years (FY2022 and earlier), whose records the triage and owner rules resolve
        fy2022 = {o["observation_id"] for o in data["observations"] if o["fiscal_year"] <= 2022
                  and o["observation_id"].startswith("OBS-LHHS-")}
        st = [v.get("human_review_status") for v in data["validations"]]
        self.assertEqual(set(st), {"", "pending", "resolved"})
        self.assertEqual(sorted(v["validation_id"] for v in data["validations"] if v["human_review_status"] == "pending"),
                         sorted(["VAL-0089", "VAL-LHHS-01816", "VAL-LHHS-01821", "VAL-LHHS-01845", "VAL-LHHS-01850",
                                 "VAL-LHHS-02006", "VAL-LHHS-03009", "VAL-LHHS-04166", "VAL-LHHS-06912"]
                                + [v["validation_id"] for v in data["validations"] if v["human_review_status"] == "pending"
                                   and (v["observation_id"] in fy2022
                                        or v["validation_id"].startswith(("VAL-TXT-", "VAL-MR-", "VAL-MRV-", "VAL-MRC-")))]))
        # the owner's text-over-table corrections (2026-10-11) close the records on the figures they replace
        # (VAL-LHHS-07178 among them) and resolve the GDM FY2020 House sum check on its replacement
        corrected = {o["observation_id"] for o in data["observations"] if o["verification_status"] == "superseded"
                     and o["observation_id"].startswith("OBS-LHHS-")} | {"OBS-LHHS-4106"}
        for v in data["validations"]:
            if v["human_review_status"] == "resolved":
                self.assertTrue(v["resolution"].strip(), v["validation_id"])
                self.assertIn(v["reviewer"], ("owner (2026-10-12)", "owner (2026-10-13)") if v["validation_id"].startswith("VAL-MRC-")
                              else ("owner (2026-10-11)",) if v["validation_id"].startswith("VAL-TXT-")
                              or (v["observation_id"] in corrected and v["reviewer"] == "owner (2026-10-11)")
                              else ("triage rules", "owner rules (2026-10-08)", "owner (2026-10-09)") if v["observation_id"] in fy2022
                              else ("owner (group approval 2026-10-08)",))
                self.assertIn(v["result"], ("fail", "flag"))
                self.assertTrue(v["resolution"])
        for sc in ("LHHS", "CJS"):
            for _, _, o in headline_observations(grid(sc)):
                self.assertTrue(all("human_review_status" in v for v in o["validation"]), o["observation_id"])
                self.assertTrue(all(v.get("resolution") for v in o["validation"] if v["human_review_status"] == "resolved"))

    def test_triage_resolutions_follow_the_proposal(self):
        import csv
        data = json.loads(S.STAGED.read_text())
        by_id = {v["validation_id"]: v for v in data["validations"]}
        with open(ROOT / "reference" / "review" / "triage_proposal.csv", newline="") as f:
            rows = [r for r in csv.DictReader(f) if r["id"].startswith("VAL-")]
        retired = {i for x in json.loads((ROOT / "reference" / "review" / "lhhs" / "retired_ids.json").read_text()).values()
                   if isinstance(x, dict) for i in x}
        for r in rows:
            if r["id"] in retired:
                continue                                 # retired by an owner decision (retired_ids.json)
            v = by_id[r["id"]]
            if v["reviewer"] == "owner (2026-10-11)" and v["resolution"].startswith("Superseded by "):
                continue                                 # its figure was replaced (owner's corrections, 2026-10-11)
            if r["proposed_disposition"] == "A":
                self.assertEqual((v["human_review_status"], v["resolution"]), ("resolved", r["reason"]), r["id"])
            elif r["family"] == "cr-estimate":
                self.assertEqual(v["human_review_status"], "resolved")
                self.assertTrue(v["resolution"].startswith("Owner decision: FY2025 was funded by a full-year CR (P.L. 119-4)"))
            else:
                self.assertEqual(v["human_review_status"], "pending", r["id"])

    def test_the_one_verified_cell_with_a_pending_record(self):
        # CJS has no flagged cell; one verified cell still has a record pending review:
        # OBS-0966 (Crime Victims Fund, FY2027 House Reported), VAL-0089's CHIMP question
        # The bill and law text check's pending records (VAL-TXT-) stand apart: on a verified cell each is an ambiguous
        # result (a flag, which counts neither way) or a disagreement with a figure a person verified (the rule leaves
        # human-verified alone)
        def pending(o, text):
            return [v["validation_id"] for v in o["validation"] if v["human_review_status"] == "pending"
                    and v["validation_id"].startswith("VAL-TXT-") == text]
        cells = [(row, key, o) for sc in ("LHHS", "CJS") for row, key, o in headline_observations(grid(sc))
                 if o["verification_status"] in ("auto-validated", "human-verified")]
        found = [(row["account"]["canonical_account_id"], key, o["observation_id"], pending(o, False))
                 for row, key, o in cells if pending(o, False)]
        self.assertEqual(found, [("ACC-DOJ-CVF", "2027|House Reported", "OBS-0966", ["VAL-0089"])])
        for row, key, o in cells:
            for v in o["validation"]:
                if v["validation_id"] in pending(o, True):
                    self.assertTrue(v["result"] == "flag" or o["verification_status"] == "human-verified",
                                    v["validation_id"])


class DraftStages(unittest.TestCase):
    def test_bill_report_reference_draft_column(self):
        data = json.loads(S.STAGED.read_text())
        self.assertEqual({r["draft"] for r in data["bill_report_refs"]}, {"TRUE", "FALSE"})
        self.assertEqual(sorted(r["reference_id"] for r in data["bill_report_refs"] if r["draft"] == "TRUE"),
                         ["BR-CJS-FY2021-SENATE", "BR-CJS-FY2022-SENATE", "BR-CJS-FY2023-SENATE", "BR-CJS-FY2024-HOUSE",
                          "BR-LHHS-FY2021-SENATE", "BR-LHHS-FY2022-SENATE", "BR-LHHS-FY2023-SENATE", "BR-LHHS-FY2024-HOUSE"])
        import openpyxl
        wb = openpyxl.load_workbook(S.reference_workbook(), read_only=True)
        head = next(wb["Bill Report Reference"].iter_rows(max_row=1, values_only=True))
        self.assertEqual(head[-2:], ("draft", "not_reported"))     # not_reported after draft (FY2020 Senate)
        self.assertEqual(sorted(r["reference_id"] for r in data["bill_report_refs"] if r.get("not_reported") == "TRUE"),
                         ["BR-LHHS-FY2020-SENATE"])

    def test_fy2024_house_reference(self):
        data = json.loads(S.STAGED.read_text())
        r = next(r for r in data["bill_report_refs"] if r["reference_id"] == "BR-LHHS-FY2024-HOUSE")
        self.assertEqual((r["lookup_key"], r["bill_id"], r["report_id"], r["draft"]),
                         ("LHHS-2024-House Reported", "H.R.5894", "JES_LHHS_H.R.5894", "TRUE"))
        self.assertEqual(r["bill_url"], "https://www.govinfo.gov/content/pkg/BILLS-118hr5894ih/pdf/BILLS-118hr5894ih.pdf")
        doc = next(d for d in data["source_docs"] if d["document_id"] == "SRC-EXPL-LHHS-FY2024-HOUSE")
        self.assertEqual(r["report_jes_url"], doc["url_or_identifier"])
        self.assertIn("408e33889225be0225d2c1fff15e095f5532df9d5b8fd8dd1c0125dbf0f8ff8d", doc["notes"])
        self.assertEqual((doc["stage"], doc["fiscal_year"], doc["document_type"]), ("House Reported", 2024, "explanatory_statement"))


class DownloadWorkbook(unittest.TestCase):
    def test_published_where_the_link_points(self):
        self.assertTrue((DOCS / "Approps_Pilot_Schema_Loaded.xlsx").is_file())
        self.assertEqual(E.DOWNLOAD, "Approps_Pilot_Schema_Loaded.xlsx")

    def test_rebuilding_twice_gives_a_byte_identical_file(self):
        with tempfile.TemporaryDirectory() as d:
            files = []
            for i in (1, 2):
                built = S.build_workbook(Path(d) / f"build{i}.xlsx")       # a fresh build each time
                out = Path(d) / f"published{i}.xlsx"
                E.publish_workbook(built, out, E.committed_date(S.STAGED))
                files.append(out.read_bytes())
            self.assertNotEqual((Path(d) / "build1.xlsx").read_bytes(), (Path(d) / "build2.xlsx").read_bytes(),
                                "openpyxl stamps the build time -- the publish step is what makes it stable")
            # the determinism check: two builds in the same environment are the same file. docs/'s copy is
            # written only by CI's export job on main; tests/test_static.py Export.test_committed_docs_are_current
            # holds its content (every cell's value and formula) to the data
            self.assertEqual(files[0], files[1])


class ReviewList(unittest.TestCase):
    def rows(self):
        import csv
        with open(ROOT / "reference" / "review" / "review_list.csv", newline="") as f:
            return list(csv.DictReader(f))

    def test_aha_link_item_removed(self):
        self.assertFalse([r for r in self.rows() if r["id"] == "SRC-CJ-AHA-FY2026"])

    def test_page_check_sheet(self):
        import csv
        with open(ROOT / "reference" / "review" / "page_check_sheet.csv", newline="") as f:
            sheet = list(csv.DictReader(f))
        page_items = [r for r in self.rows() if r["reason"] == "page not confirmed by text search"]
        self.assertEqual((len(page_items), len(sheet)), (140, 129))      # less the 11 derived figures (resolved)
        self.assertEqual(sum(bool(r["nearby_page"]) for r in sheet), 13)
        data = json.loads(S.STAGED.read_text())
        obs = {o["observation_id"]: o for o in data["observations"]}
        docs = {d["document_id"]: d["url_or_identifier"] for d in data["source_docs"]}
        import re
        for r in sheet:
            o = obs[r["observation_id"]]
            first = min(int(x) for x in re.findall(r"\d+", o["source_page"]))
            self.assertEqual(r["pdf_link"], f"{docs[o['source_document_id']]}#page={first}")
            self.assertTrue(r["pdf_link"].split("#")[0].endswith(".pdf"))
            self.assertEqual((int(r["fiscal_year"]), r["stage"]), (o["fiscal_year"], o["stage"]))
            if r["nearby_page"]:
                near = re.findall(r"\d+", r["nearby_page"])[0]
                self.assertTrue(r["nearby_link"].endswith(f"#page={near}"))
        page = (DOCS / "review" / "page-checks.html").read_text()
        self.assertEqual(page.count('data-testid="page-check-row"'), 129)
        self.assertNotIn("api_key", page)

    def test_page_checks_read(self):
        # every page check read again (page_check_read.py): 70 LHHS pages confirmed, 58 CJS ranges narrowed
        # to the one page that prints the figure; OBS-0771 (not printed as a line there) settled by the owner:
        # printed in S.Rept. 118-198 p.227
        import csv
        import re
        from collections import Counter
        with open(ROOT / "reference" / "review" / "page_check_results.csv", newline="") as f:
            res = list(csv.DictReader(f))
        self.assertEqual(Counter(r["result"] for r in res), {"confirmed": 70, "narrowed": 58, "resolved_by_owner": 1})
        self.assertEqual([r["observation_id"] for r in res if r["result"] == "resolved_by_owner"], ["OBS-0771"])
        self.assertTrue((ROOT / "reference" / "review" / "page_checks" / "OBS-0771-CRPT-118hrpt582-p248.png").is_file())
        data = json.loads(S.STAGED.read_text())
        obs = {o["observation_id"]: o for o in data["observations"]}
        for r in res:
            o = obs[r["observation_id"]]
            if r["result"] == "resolved_by_owner":
                self.assertEqual((r["human_review_status"], r["reviewer"], r["resolution"]),
                                 ("resolved", "owner 2026-10-08", "printed in S.Rept. 118-198 p.227"))
                self.assertEqual((o["source_document_id"], o["source_page"], o["extraction_method"], o["amount"]),
                                 ("SRC-CRPT-118SRPT198", "227", "human-entered", 420_000_000))
                self.assertTrue(o["source_table_or_section"].endswith(
                    "-- emergency piece of the FY2025 request; printed in the Senate report (p.227); not printed as "
                    "its own line in H.Rept. 118-582"))
                v = [v for v in data["validations"] if v["observation_id"] == r["observation_id"]
                     and v["rule_applied"] == "cross_document"]
                self.assertEqual([(x["validation_id"], x["result"]) for x in v], [("VAL-0243", "pass")])
                self.assertIn("8,045,320 = base 7,519,320 (OBS-0695) + defense 106,000 (OBS-0735) + 420,000",
                              v[0]["expected_result"])
                continue
            self.assertEqual((r["human_review_status"], r["reviewer"], r["resolution"]),
                             ("resolved", "page image check 2026-10-08", f"figure seen on PDF p.{r['page_seen']}"))
            self.assertEqual(o["source_page"], r["page_seen"])
            if r["result"] == "narrowed":
                lo, hi = (int(x) for x in re.findall(r"\d+", r["cited_page"]))
                self.assertTrue(lo <= int(r["page_seen"]) <= hi)
                self.assertTrue(o["source_table_or_section"].endswith(
                    f" -- narrowed from {r['cited_page']} by {'page image' if r['method'] == 'page image' else 'page text'}"))
        # the rows a person settled by reading (25 page images in all, under the 80 allowed; one row by its OCR text)
        reads = (ROOT / "reference" / "review" / "page_check_image_reads.jsonl").read_text().splitlines()
        self.assertEqual(len(reads), 26)

    def test_cbo_observations_tagged_interim(self):
        data = json.loads(S.STAGED.read_text())
        cbo = sorted(o["observation_id"] for o in data["observations"] if o["source_document_id"] == "SRC-CBO-HR8845-FY2027")
        # 18 of the 22 were re-sourced to H.Rept. 119-652's tables (2026-10-11); these 4 the tables don't print
        self.assertEqual(cbo, ["OBS-0966", "OBS-0967", "OBS-0970", "OBS-0983"])
        tagged = sorted(r["id"] for r in self.rows()
                        if r["reason"] == "interim source: re-source to H.Rpt. 119-652 when the rest of CJS is ingested")
        self.assertEqual(tagged, cbo)


if __name__ == "__main__":
    unittest.main()


def grid_rows(g):
    return g["rows"] + [t["total"] for t in g.get("titles", []) if t.get("total")] + \
        ([g["bill_total"]] if g.get("bill_total") else [])


class AccountingLayer(unittest.TestCase):
    """The accounting layer of the site data (UI-0, site_accounting.py): every published view, component, earmark,
    defense portion and attribute is what data/staged.json and measure_reconcile hold."""

    @classmethod
    def setUpClass(cls):
        import sys
        sys.path.insert(0, str(ROOT / "scripts"))
        import accounting as AC
        import measure_reconcile as M
        cls.AC = AC
        import tag_components as TC
        cls.TC = TC
        cls.data = json.loads(S.STAGED.read_text())
        cls.obs = {o["observation_id"]: o for o in cls.data["observations"]}
        cls.comps = {c["component_id"]: c for c in cls.data["observation_components"]}
        cls.by_obs = AC.components_by_observation(cls.data)
        cls.vals = {v["validation_id"]: v for v in cls.data["validations"]}
        records, outcomes = M.run(cls.data)
        cls.records = {r["validation_id"]: r for r in records}
        cls.outcome = {o["observation_id"]: o for o in outcomes if o["kind"] == "view"}
        cls.grids = {sc: grid(sc) for sc in json.loads((DOCS / "data" / "subcommittees.json").read_text())["subcommittees"]}

    def cells(self):
        for sc, g in self.grids.items():
            for row in grid_rows(g):
                for key, cell in row["cells"].items():
                    yield sc, row, key, cell

    def headline(self, cell):
        # the budget authority line with no component; a rescission-only cell's rescission line (D-7 cleanup)
        line = next((l for l in cell["lines"] if l["amount_type"] == "budget authority" and not l["component"]),
                    None)
        if line is None:
            import site_accounting as SA
            line = SA.rescission_line(cell)
        return self.obs[line["observations"][0]["observation_id"]]

    def test_every_view_is_measure_reconciles_record(self):
        n = 0
        for sc, row, key, cell in self.cells():
            if "views" not in cell:
                continue
            h = self.headline(cell)
            self.assertEqual(set(cell["views"]), set(self.AC.BASES), key)
            for basis, entries in cell["views"].items():
                self.assertIn(len(entries), (1, 2))
                self.assertEqual(entries[0]["timing"], None if basis == "printed" else "available")
                if len(entries) == 2:
                    self.assertEqual(entries[1]["timing"], "enacted")
                for e in entries:
                    self.assertIn(e["status"], ("checked", "computed", "matches", "excluded", "not_computable"))
                    self.assertEqual(e["status"] == "matches", bool(row["rollup"] and e.get("anchor")
                                                                    and not e.get("remainder")), e)
                    self.assertTrue(e["reason"])
                    if e["status"] in ("excluded", "not_computable"):
                        self.assertIsNone(e["amount"])
                        continue
                    self.assertIsInstance(e["amount"], int)
                    if basis == "printed":
                        self.assertEqual(e["amount"], h["amount"])
                        continue
                    if e["status"] == "checked":
                        # a printed total confirms it: measure_reconcile passed that printed view line, and the
                        # record is in data/staged.json
                        p = self.obs[e["printed_observation_id"]]
                        self.assertEqual(e["amount"], p["amount"])
                        self.assertEqual((p["canonical_account_id"], f"{p['fiscal_year']}|{p['stage']}"),
                                         (row["account"]["canonical_account_id"], key))
                        self.assertEqual(self.outcome[p["observation_id"]]["result"], "pass")
                        self.assertEqual(self.records[e["validation_id"]]["result"], "pass")
                        self.assertEqual(self.vals[e["validation_id"]]["observation_id"], p["observation_id"])
                        n += 1
                    if e.get("anchor"):
                        # a total built down from its printed line (D-8): the line is the cell's own, and the figure
                        p = self.obs[e["anchor"]["observation_id"]]
                        self.assertEqual(e["amount"], p["amount"])
                        self.assertEqual((p["canonical_account_id"], f"{p['fiscal_year']}|{p['stage']}",
                                          p["source_document_id"]),
                                         (row["account"]["canonical_account_id"], key, h["source_document_id"]))
                        n += 1
                    if "component_ids" in e:
                        # an account's view: the headline less the included components it leaves out, plus the
                        # components beside it it counts -- from the published components alone
                        used = [self.comps[c] for c in e["component_ids"]]
                        self.assertTrue(all(c["observation_id"] == h["observation_id"] for c in used))
                        got = h["amount"] + sum(c["amount"] if c["in_headline"] != "TRUE" else -c["amount"]
                                                for c in used)
                        self.assertEqual(e["amount"], got, (key, basis, e["timing"]))
                    else:
                        # a total's view: built down from its anchor (test_totals_build_down_from_their_anchor)
                        self.assertTrue(row["rollup"], (key, basis))
                        self.assertIsInstance(e["accounts_behind"], int)
        self.assertGreater(n, 150)

    def test_mixed_without_mandatory_portion_counts_as_discretionary(self):
        # ACCOUNTING_RULES: a mixed account counts its headline less its mandatory_portion; with none tagged, its
        # headline (owner, 2026-10-12)
        attrs = self.AC.attribute_index(self.data)
        n = 0
        for sc, row, key, cell in self.cells():
            if row["rollup"] or "views" not in cell:
                continue
            fy, stage = key.split("|", 1)
            rec = self.AC.attribute_value(attrs, row["account"]["canonical_account_id"], int(fy), "budget_class", stage)
            h = self.headline(cell)
            if not rec or rec["value"] != "mixed" or any(c["component_type"] == "mandatory_portion"
                                                         for c in self.by_obs.get(h["observation_id"], [])):
                continue
            for basis in ("discretionary", "discretionary_designated"):
                for e in cell["views"][basis]:
                    self.assertNotEqual(e["status"], "excluded")
                    if e["status"] in ("computed", "checked"):
                        self.assertTrue(e["reason"].startswith("mixed, no mandatory portion tagged: counted as "
                                                               "discretionary (ACCOUNTING_RULES)"), e["reason"])
                        n += 1
                    else:
                        self.assertNotIn("mixed", e["reason"])
        self.assertGreater(n, 500)

    # ------------------------------------------------------------ D-6: totals from their accounts
    def lines_behind(self, sc, row, key=None, doc=None):
        """A total's lines, worked out here from data/staged.json (owner, D-6): an agency total's accounts; a title
        or bill total's accounts plus each agency held only as its total. None for an agency held only as its
        total (it is its own line). With a cell and its table (D-9): an account that table prints under another
        agency's total moves to it (total_members_printed.csv)."""
        import csv
        import measure_reconcile as M
        acc = {a["canonical_account_id"]: a for a in self.data["accounts"]}
        a = acc[row["account"]["canonical_account_id"]]
        lone = [x["canonical_account_id"] for x in self.data["accounts"] if x.get("total_scope") == "agency"
                and x["subcommittee"] == sc and not M.members(self.data, x)]
        if a["total_scope"] == "agency":
            ms = M.members(self.data, a)
            if key and ms:
                with (ROOT / "reference" / "review" / "accounting" / "total_members_printed.csv").open() as f:
                    moved = [r for r in csv.DictReader(f) if f"{r['fiscal_year']}|{r['stage']}" == key
                             and r["document"] == doc]
                aid = a["canonical_account_id"]
                ms = [m for m in ms if not any(r["account"] == m and r["held_under"] == aid for r in moved)] + [
                    r["account"] for r in moved if r["printed_under"] == aid and r["account"] not in ms]
            return ms or None
        if a["total_scope"] == "title":
            return M.title_members(self.data, a) + [m for m in lone if (acc[m].get("title") or "") == (a.get("title") or "")]
        return [x["canonical_account_id"] for x in self.data["accounts"] if x["subcommittee"] == sc
                and not x.get("total_scope") and not x.get("parent_account_id")] + lone

    def test_totals_build_down_from_their_anchor(self):
        # D-8 (reference/specs/anchored-totals.md): every entry of a total but its As printed is built down from its
        # printed anchor -- matches (the anchor, every step itemised), computed against an anchor (the anchor, with
        # one 'not itemised' remainder naming the parts not known), computed with no anchor (As printed + the
        # adjustments, every part known), excluded, or not computable with a cause. Never a sum chosen to match
        import site_accounting as SA
        rows = {sc: {r["account"]["canonical_account_id"]: r for r in grid_rows(g)} for sc, g in self.grids.items()}
        seen = collections.Counter()
        for sc, row, key, cell in self.cells():
            if not row["rollup"] or "views" not in cell:
                continue
            lines = self.lines_behind(sc, row, key, self.headline(cell)["source_document_id"]
                                      if "views" in cell else None)
            if lines is None:
                self.assertNotIn("accounts_behind", cell["views"]["all_appropriations"][0])
                continue
            inlife = [rows[sc][m]["cells"][key] for m in lines if not rows[sc][m]["cells"][key].get("outside_life")]
            held = [c for c in inlife if "views" in c]
            printed = cell["views"]["printed"][0]["amount"]
            whole = sum(c["views"]["printed"][0]["amount"] for c in held) == printed
            for field in ("views",) + tuple(SA.EXCLUSION_FIELDS.values()):
                for basis, entries in cell[field].items():
                    if field == "views" and basis == "printed":
                        self.assertEqual(entries[0]["status"], "checked")
                        continue
                    for e in entries:
                        where = (row["account"]["canonical_account_id"], key, field, basis, e["timing"])
                        self.assertIsInstance(e["accounts_behind"], int)
                        if e["status"] == "not_computable":
                            self.assertTrue(e["reason"].startswith("not computable"), e["reason"])
                            self.assertIn(e["cause"], SA.CAUSES)
                            self.assertNotIn("anchor", e)
                            seen["not_computable"] += 1
                            continue
                        self.assertNotIn("cause", e)
                        mine = [SA.pick(c[field][basis], e["timing"]) for c in held]
                        if e["status"] == "excluded":
                            self.assertTrue(all(x["status"] == "excluded" for x in mine))
                            continue
                        if e.get("anchor"):
                            a = self.obs[e["anchor"]["observation_id"]]
                            self.assertEqual(a["source_document_id"], self.headline(cell)["source_document_id"])
                            self.assertIn(basis, [b for b, _ in SA.PRINTED_BASIS[a["component"]]], where)
                            self.assertEqual(e["anchor"]["label"], self.TC.printed(a))
                            rem = e.get("remainder")
                            if e["status"] == "matches":
                                self.assertIsNone(rem)
                                if field == "views":
                                    self.assertEqual(e["amount"], a["amount"], where)
                            else:
                                self.assertEqual(e["status"], "computed")
                                if field == "views":
                                    self.assertEqual(e["amount"], a["amount"], where)
                                    self.assertTrue(rem["amount"])
                                    last = e["adjustments"][-1]
                                    self.assertEqual((last["type"], last["amount"]), ("not_itemised", abs(rem["amount"])))
                                    self.assertEqual(last["sign"], "+" if rem["amount"] > 0 else "-")
                                seen["remainder"] += 1
                            seen[e["status"]] += 1
                            continue
                        # no anchor: every part known, As printed + the adjustments
                        self.assertEqual(e["status"], "computed", where)
                        self.assertNotIn("remainder", e)
                        self.assertTrue(all(x["status"] in SA.VALUED + ("excluded",) for x in mine), where)
                        if whole and field == "views" and not any(a["type"] == "held_line_not_counted"
                                                                  for a in e["adjustments"]):
                            # the accounts add up to the total as printed: the build-up is their sum
                            self.assertEqual(e["amount"], sum(x["amount"] for x in mine if x["status"] != "excluded"),
                                             where)
                            seen["sum"] += 1
                        seen["computed"] += 1
        self.assertGreater(seen["computed"], 3000)
        self.assertGreater(seen["matches"], 500)
        self.assertGreater(seen["remainder"], 50)
        self.assertGreater(seen["sum"], 900)

    def test_gap_totals_are_listed_for_review(self):
        # a total whose accounts as printed don't add up to it (money in no account held) is in
        # reference/review/accounting/total_gaps.csv; the money stays in the build-up (D-8): against an anchor it is
        # in the remainder, with no anchor the entry is not computable (cause not_held) -- never assumed to hold
        # nothing a basis needs
        import csv
        listed = {(r["total"], f"{r['fiscal_year']}|{r['stage']}"): r for r in
                  csv.DictReader((ROOT / "reference" / "review" / "accounting" / "total_gaps.csv").open())}
        found = set()
        for sc, row, key, cell in self.cells():
            if not row["rollup"] or "views" not in cell:
                continue
            aid = row["account"]["canonical_account_id"]
            for basis, entries in cell["views"].items():
                for e in entries:
                    parts = [p for p in (e.get("remainder") or {}).get("accounts", []) if p["account"] is None]
                    if e.get("cause") == "not_held" or parts:
                        found.add((aid, key))
                    if parts and (aid, key) in listed:
                        self.assertEqual(parts[0]["amount"], int(listed[(aid, key)]["short"]), (aid, key))
        gaps = {k for k, r in listed.items() if r["published"] == "no"}
        self.assertLessEqual(gaps, found)
        # the rest found: cells with a line with no figure (not listed: the gap is not the held lines')
        for aid, key in found - gaps:
            cell = next(c for sc, r, k, c in self.cells() if r["account"]["canonical_account_id"] == aid and k == key)
            self.assertTrue(any(e.get("cause") in ("missing_figure", "no_anchor") or any(
                p["cause"] == "missing_figure" for p in (e.get("remainder") or {}).get("accounts", []))
                for es in cell["views"].values() for e in es), (aid, key))
        self.assertEqual({r["published"] for r in listed.values()}, {"yes", "no"})
        for k, r in listed.items():
            self.assertEqual(r["published"] == "yes", r["lines_printed_inside"] != "" and int(r["short"]) == 0, k)
        self.assertEqual(listed[("ACC-HHS-CMS-TOTAL", "2021|House Reported")]["short"], "1000")
        self.assertGreater(len(gaps), 20)      # D-9 closed some (29 left): the list is still more than a handful

    # ------------------------------------------------------------ D-6 round 3: held lines need position evidence
    def total_cell(self, aid, fy, stage):
        import site_accounting as SA
        rows = {r["account"]["canonical_account_id"]: r for g in self.grids.values() for r in grid_rows(g)}
        layer = SA.Layer(self.data)
        layer.rows = rows
        a = layer.acc[aid]
        cell = rows[aid]["cells"][f"{fy}|{stage}"]
        h = layer.obs[cell["views"]["printed"][0]["printed_observation_id"]]
        return layer, a, h, cell

    def test_a_held_line_outside_the_total_is_never_counted(self):
        # owner, 2026-10-14: a held line counts only because the table prints it inside that total (its position or
        # the table's printed arithmetic: total_lines_inside.csv), never because it makes the sum match
        import csv
        import site_accounting as SA
        evidence = list(csv.DictReader((ROOT / "reference" / "review" / "accounting" / "total_lines_inside.csv").open()))
        inside = {(r["total_observation_id"], r["line_observation_id"]) for r in evidence if r["inside"] == "yes"}
        # a line from another document than the total's table is outside, except an advance tagged from the bill or
        # law text (a component) that the table prints inside the total (owner, 2026-10-15, D-9)
        self.assertTrue(all(r["line_document"] == r["document"] or r["line_observation_id"].startswith("CMP-")
                            for r in evidence if r["inside"] == "yes"))
        self.assertTrue(all(r["where"] and r["evidence"] for r in evidence if r["inside"] == "yes"))
        cases = [
            # printed outside: another document's line (P.L. 114-254's MCH line beside H.Rept. 115-244's total), closing
            # the gap exactly. (NSF FY2025's request emergency counts since D-9: S.Rept. 118-198's request column prints
            # it inside the same request totals, owner 2026-10-15)
            ("ACC-HHS-HRSA-TOTAL", 2017, "Enacted", None),
            # printed inside, but with its evidence taken away: not counted either
            ("ACC-NASA-TOTAL", 2024, "Enacted", "drop"),
        ]
        for aid, fy, stage, how in cases:
            with self.subTest(total=aid, fy=fy, stage=stage):
                layer, a, h, cell = self.total_cell(aid, fy, stage)
                held = SA.held_lines(layer.behind(a, f"{fy}|{stage}")[0])
                gap = h["amount"] - sum(mc["views"]["printed"][0]["amount"] for _, mc in layer.behind(a, f"{fy}|{stage}")[0])
                closing = [x for x in held if x["amount"] == gap] or held
                self.assertEqual(sum(x["amount"] for x in closing), gap)       # the arithmetic would match
                pairs = {(h["observation_id"], x["observation_id"]) for x in closing}
                if how == "drop":
                    self.assertTrue(pairs <= inside)
                else:
                    self.assertFalse(pairs & inside)
                layer.inside = inside - pairs
                out = layer.total_layer(a, fy, stage, h, cell)
                e = out["views"]["all_appropriations"][0]
                self.assertEqual(e["status"], "not_computable", e)
                self.assertIn("is in no account held", e["reason"])
                # the same lines with evidence would be counted: the arithmetic alone is what the rule refuses
                layer.inside = inside | pairs
                self.assertEqual(layer.total_layer(a, fy, stage, h, cell)["views"]["printed"][0]["status"], "checked")
                self.assertIn(layer.total_layer(a, fy, stage, h, cell)["views_excluding_earmarks"]["printed"][0]["status"],
                              ("computed", "checked"))
        # a line the evidence marks outside (NASA FY2025 Enacted's 'Other appropriations' $740,200,000, printed on
        # p.261 after the NASA total) is not counted
        out_rows = [r for r in evidence if r["inside"] == "no"]
        self.assertTrue(out_rows)
        layer, a, h, cell = self.total_cell("ACC-NASA-TOTAL", 2025, "Enacted")
        for r in out_rows:
            self.assertNotIn((r["total_observation_id"], r["line_observation_id"]), layer.inside)


    def test_a_line_counts_in_the_total_the_table_prints_it_under(self):
        # owner, 2026-10-15 (D-9): ARPA-H, held under NIH, is printed under the Office of the Secretary in H.Rept.
        # 117-403 (p.836): in those two cells it is behind the OS total and not NIH's, and nowhere else; each row's
        # table is the total's own; NIH's proposed NIRSQ, held on the NIH total, counts only with its evidence row
        import csv
        import site_accounting as SA
        with SA.MEMBERS_FILE.open() as f:
            moved = list(csv.DictReader(f))
        self.assertTrue(moved)
        totals = {(o["canonical_account_id"], o["fiscal_year"], o["stage"], o["source_document_id"])
                  for o in self.data["observations"] if not o["component"] and o["verification_status"] != "superseded"}
        for r in moved:
            self.assertIn((r["printed_under"], int(r["fiscal_year"]), r["stage"], r["document"]), totals, r)
            self.assertTrue(r["where"] and r["evidence"], r)
        arpa = "ACC-HHS-NIH-ARPA-H"
        for fy, stage in ((2022, "Enacted"), (2023, "House Reported")):
            layer, a, h, cell = self.total_cell("ACC-HHS-OS-TOTAL", fy, stage)
            self.assertIn(arpa, [m for m, _ in layer.behind(a, f"{fy}|{stage}", h["source_document_id"])[0]])
            self.assertNotIn(arpa, [m for m, _ in layer.behind(a, f"{fy}|{stage}")[0]])      # no table: as held
            layer, a, h, cell = self.total_cell("ACC-HHS-NIH-TOTAL", fy, stage)
            self.assertNotIn(arpa, [m for m, _ in layer.behind(a, f"{fy}|{stage}", h["source_document_id"])[0]])
        layer, a, h, cell = self.total_cell("ACC-HHS-NIH-TOTAL", 2023, "Senate Reported")
        self.assertIn(arpa, [m for m, _ in layer.behind(a, "2023|Senate Reported", h["source_document_id"])[0]])
        layer, a, h, cell = self.total_cell("ACC-HHS-OS-TOTAL", 2022, "Enacted")
        self.assertEqual(layer.total_layer(a, 2022, "Enacted", h, cell)["views"]["printed"][0]["status"], "checked")
        self.assertFalse([g for g in layer.gaps if g["total"] == "ACC-HHS-OS-TOTAL"])
        # NIRSQ FY2018: closed by its evidence row; without it, the same money is in no account held
        layer, a, h, cell = self.total_cell("ACC-HHS-NIH-TOTAL", 2018, "President's Budget")
        pair = (h["observation_id"], "OBS-LHHS-3504")
        self.assertIn(pair, layer.inside)
        self.assertNotIn(SA.NOT_HELD, [u["cause"] for u in layer.line_parts(
            dict(self.parts(layer, a, h, cell)), "views", "program_level", "enacted")[3]])
        layer.inside = layer.inside - {pair}
        self.assertIn(SA.NOT_HELD, [u["cause"] for u in layer.line_parts(
            dict(self.parts(layer, a, h, cell)), "views", "program_level", "enacted")[3]])

    def parts(self, layer, a, h, cell):
        # the parts total_layer builds, read back through the gaps it records (one total cell)
        layer.gaps = []
        captured = {}
        orig = layer.line_parts
        layer.line_parts = lambda parts, *x: captured.setdefault("p", parts) and orig(parts, *x)
        try:
            layer.total_layer(a, int(h["fiscal_year"]), h["stage"], h, cell)
        finally:
            layer.line_parts = orig
        return captured["p"]

    def test_every_component_is_in_staged(self):
        n = 0
        for sc, row, key, cell in self.cells():
            if "components" not in cell:
                continue
            h = self.headline(cell)
            self.assertEqual(sorted(c["component_id"] for c in cell["components"]),
                             sorted(c["component_id"] for c in self.by_obs.get(h["observation_id"], [])))
            for c in cell["components"]:
                s = self.comps[c["component_id"]]
                self.assertEqual((c["amount"], c["type"], c["in_headline"], c["source_document_id"],
                                  c["within_component_id"] or ""),
                                 (s["amount"], s["component_type"], s["in_headline"] == "TRUE",
                                  s["source_document_id"], s.get("within_component_id") or ""))
                n += 1
            for field, ctype in (("earmark", "earmark"), ("defense_portion", "defense_portion")):
                tagged = [c for c in cell["components"] if c["type"] == ctype]
                if not tagged:
                    self.assertNotIn(field, cell)            # none tagged is not zero
                    continue
                self.assertEqual(sorted(cell[field]["component_ids"]),
                                 sorted(c["component_id"] for c in tagged if not c["within_component_id"]
                                        or c["within_component_id"] not in {x["component_id"] for x in tagged}))
                self.assertEqual(cell[field]["amount"],
                                 sum(self.comps[i]["amount"] for i in cell[field]["component_ids"]))
        self.assertEqual(n, sum(1 for c in self.data["observation_components"]
                                if self.obs[c["observation_id"]]["canonical_account_id"] in
                                {r["account"]["canonical_account_id"] for g in self.grids.values() for r in grid_rows(g)}))

    def test_no_view_where_the_base_is_negative(self):
        for sc, row, key, cell in self.cells():
            if "views" not in cell:
                continue
            h = self.headline(cell)
            base = self.AC.split(h["amount"], self.by_obs.get(h["observation_id"], []))[0]
            if h["amount_type"] == "rescission":
                # a rescission line is negative as printed (its sign is its effect on budget authority): its views
                # are the rescission, never a negative base (D-7 cleanup)
                self.assertEqual(cell["views"]["printed"][0]["amount"], h["amount"])
                continue
            if base < 0:
                for basis, entries in cell["views"].items():
                    if basis == "printed" and h["amount"] < 0 and re.search(r"-(TITLE-[IVX]+|BILL)-TOTAL$",
                                                                             row["account"]["canonical_account_id"]):
                        # a total printed negative (CJS Title V, General Provisions: net of its rescissions, D-7)
                        # is shown as printed, and only as printed
                        self.assertEqual([(e["amount"], e["status"]) for e in entries], [(h["amount"], "checked")])
                        continue
                    self.assertTrue(all(e["amount"] is None and e["status"] == "not_computable" for e in entries))

    def test_views_only_on_cells_with_a_printed_figure(self):
        for sc, row, key, cell in self.cells():
            line = next((l for l in cell["lines"] if l["amount_type"] == "budget authority" and not l["component"]),
                        None)
            if line is None:
                # a rescission-only cell (the HHS general provisions' rescissions): its rescission line (D-7 cleanup)
                line = next((l for l in cell["lines"] if l["amount_type"] == "rescission" and not l["component"]),
                            None) if not any(l["amount_type"] == "budget authority" for l in cell["lines"]) else None
            has = bool(line and line["state"] in ("value", "not_funded") and line["observations"])
            self.assertEqual("views" in cell, has, (row["account"]["canonical_account_id"], key))

    def test_status_counts_per_basis(self):
        import site_accounting as SA
        for sc, g in self.grids.items():
            got = SA.status_counts(g)
            self.assertEqual(got, GRID_COUNTS["view_status_counts"][sc], sc)
            # every value cell has a printed view, and a mandatory account is excluded only from the
            # discretionary bases
            self.assertTrue(all(k.endswith("checked") for k in got["printed"]))
            for basis in ("all_appropriations", "program_level"):
                self.assertFalse([k for k in got[basis] if k.endswith("excluded")], basis)

    # ------------------------------------------------------------ UI-0b: excluding earmarks, nondefense
    def exclusion_entries(self, cell):
        import site_accounting as SA
        for exclusion, field in SA.EXCLUSION_FIELDS.items():
            for basis, entries in cell.get(field, {}).items():
                yield exclusion, field, basis, entries

    def test_exclusion_views_on_every_view_cell(self):
        import site_accounting as SA
        for sc, row, key, cell in self.cells():
            for field in SA.EXCLUSION_FIELDS.values():
                self.assertEqual(field in cell, "views" in cell, (key, field))
                if field not in cell:
                    continue
                self.assertEqual(set(cell[field]), set(self.AC.BASES), key)
                for basis, entries in cell[field].items():
                    self.assertIn(len(entries), (1, 2))
                    self.assertEqual(entries[0]["timing"], None if basis == "printed" else "available")
                    if len(entries) == 2:
                        self.assertEqual(entries[1]["timing"], "enacted")
                        self.assertNotEqual((entries[0]["amount"], entries[0]["status"]),
                                            (entries[1]["amount"], entries[1]["status"]))
                    for e in entries:
                        self.assertIn(e["status"], ("checked", "computed", "matches", "excluded", "not_computable"))
                        self.assertTrue(e["reason"])
                        self.assertEqual(e["amount"] is None, e["status"] in ("excluded", "not_computable"))
                        if e["status"] == "not_computable":
                            self.assertIn(e["cause"], SA.CAUSES)

    def test_each_account_exclusion_view_is_accountings_result(self):
        # every published entry of an account cell equals accounting.excluded_amount for that cell, basis and timing
        import tag_components as TC
        attrs = self.AC.attribute_index(self.data)
        n = 0
        for sc, row, key, cell in self.cells():
            if row["rollup"] or "views" not in cell:
                continue
            aid = row["account"]["canonical_account_id"]
            fy, stage = key.split("|", 1)
            if (sc, int(fy), stage) not in TC.auto_cells():
                continue
            h = self.headline(cell)
            cs = self.by_obs.get(h["observation_id"], [])
            # a rescission line is negative as printed, not a negative base (D-7 cleanup)
            if self.AC.split(h["amount"], cs)[0] < 0 and h["amount_type"] != "rescission":
                for _, field, basis, entries in self.exclusion_entries(cell):
                    self.assertEqual(entries, cell["views"][basis])
                continue
            rec = self.AC.attribute_value(attrs, aid, int(fy), "budget_class", stage)
            for exclusion, field, basis, entries in self.exclusion_entries(cell):
                timings = [None] if basis == "printed" else self.AC.TIMINGS
                want = [self.AC.excluded_amount(h["amount"], cs, basis, t, rec["value"] if rec else None, exclusion)
                        for t in timings]
                got = {e["timing"]: e for e in entries}
                for t, w in zip(timings, want):
                    e = got.get(t) or got[timings[0]]          # no 'enacted' entry: 'available' holds for both
                    if w[0] != "value":
                        self.assertEqual((e["status"], e["amount"], e["reason"]), (w[0], None, w[1]))
                        continue
                    self.assertEqual(e["status"], "computed", (aid, key, field, basis))   # an account's is never checked
                    self.assertEqual(e["amount"], w[1], (aid, key, field, basis, t))
                    if e["timing"] != t:
                        continue        # the 'available' entry holds for 'enacted' too: same amount and status
                    self.assertEqual(e["component_ids"], sorted(c["component_id"] for c in w[2]))
                    # and from the published components alone: the headline less those inside, plus those beside
                    used = [self.comps[c] for c in e["component_ids"]]
                    self.assertEqual(e["amount"], h["amount"] + sum(c["amount"] if c["in_headline"] != "TRUE"
                                                                    else -c["amount"] for c in used))
                    import site_accounting as SA
                    taken = [c for c in used if c["component_type"] in SA.EXCLUSION_TYPES[exclusion]]
                    self.assertTrue(all(c["in_headline"] == "TRUE" for c in taken))
                    n += 1
        self.assertGreater(n, 20000)

    def test_earmarks_come_out_where_tagged(self):
        # a cell with an earmark inside its headline: the excluding-earmarks printed view is the headline less it
        n = 0
        for sc, row, key, cell in self.cells():
            if row["rollup"] or "earmark" not in cell or cell["views_excluding_earmarks"]["printed"][0]["status"] != "computed":
                continue
            h = self.headline(cell)
            inside = sum(self.comps[i]["amount"] for i in cell["earmark"]["component_ids"]
                         if self.comps[i]["in_headline"] == "TRUE")
            self.assertEqual(cell["views_excluding_earmarks"]["printed"][0]["amount"], h["amount"] - inside)
            n += 1
        self.assertGreater(n, 20)

    def test_untagged_stages_are_never_computed(self):
        # outside the auto-tagged cells (FY2027 but its House stage (D-1), ...) untagged is never 'no earmarks' or
        # 'no defense portion': every entry is not_computable with the coverage reason. The President's Budget is
        # read account by account (D-2): an account whose request language was not found is untagged in the same
        # way, its reason saying why; one whose language was read is covered, like an auto-tagged cell
        import site_accounting as SA
        import tag_components as TC
        seen = set()
        for sc, row, key, cell in self.cells():
            fy, stage = key.split("|", 1)
            if (sc, int(fy), stage) in TC.auto_cells():
                self.assertNotEqual(stage, "President's Budget")
                if fy == "2027":
                    self.assertEqual(stage, "House Reported")     # FY2027: only the House bills' text is held
                continue
            account = row["account"]["canonical_account_id"]
            if stage == "President's Budget" and (row["rollup"] or TC.covered(account, sc, int(fy), stage)):
                seen.add("request read" if not row["rollup"] else "request total")
                continue
            if row["rollup"] and self.lines_behind(sc, row) is not None:
                continue        # a total builds down from its anchor (test_totals_build_down_from_their_anchor)
            for exclusion, field, basis, entries in self.exclusion_entries(cell):
                if cell[field]["printed"][0]["reason"].startswith(("a request carries no earmarks",
                                                                   "not a defense function", SA.RESCISSION_NO_EARMARKS)):
                    seen.add("structural")
                    continue    # a structural rule (test_structural_rules)
                if fy == "2027":
                    seen.add("2027")
                self.assertEqual(len(entries), 1)
                e = entries[0]
                self.assertEqual((e["status"], e["amount"]), ("not_computable", None), (key, field, basis))
                # both exclusions together (D-8): a part's own untagged reason where a structural rule cleared the other
                self.assertTrue(e["reason"].startswith(tuple(SA.UNTAGGED[p] for p in SA.EXCLUSION_PARTS[exclusion])
                                                       + (SA.UNTAGGED[exclusion],)), (key, field, basis, e["reason"]))
                if stage == "President's Budget":
                    self.assertIn("request language not found", e["reason"], (key, field, basis))
                elif len(SA.EXCLUSION_PARTS[exclusion]) == 1:
                    self.assertEqual(e["reason"], SA.UNTAGGED[exclusion], (key, field, basis))
                self.assertEqual(e["cause"], "components_unknown")
                self.assertNotIn("component_ids", e)
                if row["rollup"]:
                    self.assertIsInstance(e["accounts_behind"], int)
        # the President's Budget: every excluding-earmarks view by the structural rule (D-6); the nondefense views
        # of accounts it doesn't cover keep the coverage reason
        self.assertTrue({"2027", "request read", "structural"} <= seen, seen)
        self.assertEqual(SA.UNTAGGED, {"excluding_earmarks": "earmarks not tagged for this stage",
                                       "nondefense": "defense not tagged for this stage",
                                       "nondefense_excluding_earmarks": "earmarks or defense not tagged for this stage"})

    def test_structural_rules(self):
        # D-6 (owner): a request carries no earmarks; an account outside function 050 (FY N's own Appendix volume),
        # with no defense line printed and no defense_portion tagged, holds no defense. Either way the exclusion view
        # is the cell's view entry for entry; a view that is excluded or not computable stays so
        attrs = self.AC.attribute_index(self.data)
        n = collections.Counter()
        for sc, row, key, cell in self.cells():
            if "views" not in cell or (row["rollup"] and self.lines_behind(sc, row) is not None):
                continue
            fy, stage = key.split("|", 1)
            aid = row["account"]["canonical_account_id"]
            h = self.headline(cell)
            rec = self.AC.attribute_value(attrs, aid, int(fy), "budget_subfunction", stage)
            nondefense = (rec is not None and rec["source_document_id"].startswith(f"SRC-BUDGET-APP-FY{fy}")
                          and rec["value"] != "999" and not rec["value"].startswith("05")
                          and not any(l["component"] in ("defense", "defense_inside") and l["observations"]
                                      for l in cell["lines"])
                          and not any(c["component_type"] == "defense_portion"
                                      for c in self.by_obs.get(h["observation_id"], [])))
            rescission = h["amount_type"] == "rescission" and stage != "President's Budget"
            for field, applies, why in (("views_excluding_earmarks", stage == "President's Budget" or rescission,
                                         "a rescission carries no earmarks" if rescission
                                         else "a request carries no earmarks"),
                                        ("views_nondefense", nondefense, "not a defense function (subfunction ")):
                for basis, entries in cell[field].items():
                    base = cell["views"][basis]
                    self.assertEqual(len(entries) if applies else 0, len(base) if applies else 0)
                    for e, b in zip(entries, base):
                        if not applies:
                            self.assertNotIn(why, e["reason"])
                            continue
                        self.assertEqual((e["amount"], e["timing"]), (b["amount"], b["timing"]), (aid, key, field))
                        if b["status"] in ("checked", "computed"):
                            # an agency held only as its total keeps a checked view (site_accounting.lone_checks)
                            self.assertIn(e["status"], ("computed", "checked") if row["rollup"] else ("computed",))
                            self.assertTrue(e["reason"].startswith(why), e["reason"])
                            n[field] += 1
                        else:
                            self.assertEqual((e["status"], e["reason"]), (b["status"], b["reason"]))
        self.assertGreater(n["views_excluding_earmarks"], 2000)
        self.assertGreater(n["views_nondefense"], 10000)

    def test_adjustments_turn_as_printed_into_every_entry(self):
        # D-6: every entry with an amount lists what turns As printed into it; As printed + the adjustments = the
        # amount, exactly. An account's name each component with its source; a total's group them by type, with
        # the accounts behind each
        n = collections.Counter()
        for sc, row, key, cell in self.cells():
            if "views" not in cell:
                continue
            printed = cell["views"]["printed"][0]["amount"]
            comps = {c["component_id"]: c for c in cell["components"]}
            for field in ("views", "views_excluding_earmarks", "views_nondefense", "views_nondefense_excluding_earmarks"):
                for basis, entries in cell[field].items():
                    for e in entries:
                        if e["amount"] is None:
                            self.assertNotIn("adjustments", e)
                            continue
                        adj = e["adjustments"]
                        got = printed + sum(x["amount"] if x["sign"] == "+" else -x["amount"] for x in adj)
                        self.assertEqual(got, e["amount"], (row["account"]["canonical_account_id"], key, field, basis))
                        self.assertTrue(all(x["sign"] in "+-" and x["amount"] > 0 for x in adj))
                        if e["amount"] == printed:
                            n["equal"] += 1
                            if adj:     # only offsetting adjustments leave the amount as printed
                                self.assertTrue({"+", "-"} <= {x["sign"] for x in adj})
                        for x in adj:
                            if "component_id" in x:
                                c = comps[x["component_id"]]
                                self.assertEqual((x["type"], x["source_document_id"]),
                                                 (c["type"], c["source_document_id"]))
                                self.assertEqual(x["sign"], "-" if c["in_headline"] else "+")
                                n["component"] += 1
                            else:
                                self.assertTrue(row["rollup"])
                                self.assertNotEqual(x["type"], "printed_total")      # retired (D-8)
                                if x["type"] == "held_line_not_counted":
                                    self.assertEqual(x["amount"], abs(sum(l["amount"] for l in x["lines"])))
                                elif x["type"] == "not_itemised":
                                    # the remainder against the anchor, last, cited to the anchor's page
                                    self.assertIs(x, adj[-1])
                                    self.assertEqual(x["amount"], abs(e["remainder"]["amount"]))
                                    self.assertEqual(x["printed_observation_id"], e["anchor"]["observation_id"])
                                else:
                                    self.assertEqual(x["amount"], sum(a["amount"] for a in x["accounts"]))
                                n[x["type"] if x["type"] in ("left_out_by_basis", "not_itemised") else "group"] += 1
                        n["entries"] += 1
        self.assertGreater(n["entries"], 50000)
        self.assertTrue(n["component"] and n["group"] and n["left_out_by_basis"] and n["not_itemised"] and n["equal"], n)

    def test_spot_checks(self):
        # the brief's figures: CDC Injury Prevention FY2025 request, program level: + $100,000,000 transfer in from
        # section 241; NSF R&RA FY2026 Enacted, nondefense: - $118,800,000; HRSA's FY2025 request total excluding
        # earmarks is shown
        rows = {r["account"]["canonical_account_id"]: r for g in self.grids.values() for r in grid_rows(g)}
        e = rows["ACC-HHS-CDC-INJURY-PREVENTION"]["cells"]["2025|President's Budget"]["views"]["program_level"][0]
        self.assertEqual([(x["sign"], x["amount"], x["type"], x["component_id"]) for x in e["adjustments"]],
                         [("+", 100_000_000, "transfer_in", "CMP-LHHS-1372")])
        self.assertIn("section 241", e["adjustments"][0]["label"])
        e = rows["ACC-NSF-RRA"]["cells"]["2026|Enacted"]["views_nondefense"]["all_appropriations"][0]
        self.assertEqual([(x["sign"], x["amount"], x["type"]) for x in e["adjustments"]],
                         [("-", 118_800_000, "defense_portion")])
        self.assertEqual(e["amount"], 7_057_700_000)
        # CMS FY2024 Enacted (D-6 round 2): Medicaid's advance for FY2025, printed inside the total, held beside its
        # headline; advances when enacted the sum is the printed total, when available the prior year's advance instead
        cms = rows["ACC-HHS-CMS-TOTAL"]["cells"]["2024|Enacted"]["views"]["all_appropriations"]
        self.assertEqual({e["timing"]: e["amount"] for e in cms},
                         {"available": 1_085_847_068_000, "enacted": 1_133_847_008_000})
        e = rows["ACC-HHS-HRSA-TOTAL"]["cells"]["2025|President's Budget"]["views_excluding_earmarks"]["all_appropriations"][0]
        self.assertEqual((e["status"], e["amount"]), ("computed", 8_530_564_000))
        # D-8: CDC FY2025 House (the spec's case), program level: As printed + the PPHF transfers reach the report's
        # 'Total, Centers for Disease Control, program level' (H.Rept. 118-585 p.316). D-9 part 1 entered CDC-Wide's
        # figure from H.R. 9029 p.64, so the 'available' timing is known too and both timings are one entry
        cdc = rows["ACC-HHS-CDC-TOTAL"]["cells"]["2025|House Reported"]["views"]["program_level"]
        self.assertEqual([(e["status"], e["amount"]) for e in cdc], [("matches", 7_446_058_000)])
        self.assertEqual((cdc[0]["anchor"]["document"], cdc[0]["anchor"]["page"]), ("SRC-CRPT-118HRPT585", "316"))
        self.assertEqual({x["type"] for x in cdc[0]["adjustments"]}, {"transfer_in"})

    def test_exclusion_status_counts_per_basis(self):
        import site_accounting as SA
        for sc, g in self.grids.items():
            for field in SA.EXCLUSION_FIELDS.values():
                self.assertEqual(SA.status_counts(g, field), GRID_COUNTS[f"{field}_status_counts"][sc], (sc, field))

    def test_attributes_cite_their_records(self):
        recs = {r["account_attribute_id"]: r for r in self.data["account_attributes"]}
        n = 0
        for sc, g in self.grids.items():
            for row in grid_rows(g):
                for fy, entry in row.get("attributes", {}).items():
                    if "citations" not in entry:
                        self.assertEqual(list(entry), ["request"])   # a proposed-only year: the request's values
                    for e in (entry, entry.get("request")):
                        if not e or "citations" not in e:
                            continue
                        for attr, cite in e["citations"].items():
                            r = recs[cite["account_attribute_id"]]
                            self.assertEqual((r["canonical_account_id"], str(r["fiscal_year"]), r["attribute"],
                                              r["value"], r["stage_basis"], r["source_document_id"]),
                                             (row["account"]["canonical_account_id"], fy, attr, e[attr],
                                              e["stage_basis"], cite["source_document_id"]))
                            n += 1
                        if "budget_class" in e:
                            self.assertEqual(e["provisional"],
                                             bool(recs[e["citations"]["budget_class"]["account_attribute_id"]]
                                                  .get("provisional")))
        self.assertEqual(n, len(self.data["account_attributes"]))

    def test_split_files_hold_the_same_grid(self):
        # a subcommittee past SPLIT_BYTES is published one file per fiscal year instead of whole (R1: not both);
        # test_static.py checks the years put back together are the live answer
        index = json.loads((DOCS / "data" / "index.json").read_text())
        for sc, g in self.grids.items():
            folder = DOCS / "data" / "subcommittees" / sc
            whole = DOCS / "data" / "subcommittees" / f"{sc}.json"
            size = len((json.dumps(g, sort_keys=True, ensure_ascii=False, separators=(",", ":")) + "\n").encode())
            self.assertEqual(sc in index["split_subcommittees"], size > E.SPLIT_BYTES, sc)
            self.assertEqual(folder.is_dir(), size > E.SPLIT_BYTES, sc)
            self.assertEqual(whole.exists(), size <= E.SPLIT_BYTES, sc)
            if not folder.is_dir():
                continue
            head = json.loads((folder / "index.json").read_text())
            self.assertEqual(head["fiscal_years"], g["fiscal_years"])
            self.assertEqual(sorted(p.name for p in folder.glob("*.json")),
                             sorted(["index.json"] + list(head["files"].values())))
            parts = [json.loads((folder / head["files"][str(fy)]).read_text()) for fy in g["fiscal_years"]]
            for p in parts:
                self.assertLess((folder / f"{p['fiscal_year']}.json").stat().st_size, E.SPLIT_BYTES)
            for i, row in enumerate(g["rows"]):
                cells = {}
                for p in parts:
                    self.assertEqual(p["rows"][i]["account"], row["account"])
                    self.assertTrue(all(k.startswith(f"{p['fiscal_year']}|") for k in p["rows"][i]["cells"]))
                    cells.update(p["rows"][i]["cells"])
                self.assertEqual(cells, row["cells"])
